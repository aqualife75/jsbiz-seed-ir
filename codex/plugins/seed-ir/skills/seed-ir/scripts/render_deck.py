#!/usr/bin/env python3
"""Render a reviewed Seed IR contract to editable PPTX and a content storyboard.

The storyboard is deliberately not advertised as a PowerPoint rendering. Run
render_slides.ps1 on Windows with PowerPoint and inspect every resulting slide.
Dependencies: python-pptx >= 1.0, Pillow >= 10.
"""
from __future__ import annotations

import argparse
import base64
import html
import hashlib
import importlib.util
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any

from PIL import Image, ImageFont, UnidentifiedImageError
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


class RenderError(ValueError):
    """An actionable input failure; never silently omit required deck content."""


DEFAULT_COLORS = {
    'ink': '10141B', 'paper': 'F5F2EB', 'white': 'FFFFFF',
    'accent': 'E0492E', 'gold': 'D4AF37', 'amber': 'F2A33C',
    'teal': '4FD1C5', 'teal_dark': '177B72', 'muted': '5C6572',
    'muted_dark': 'B3BAC4', 'panel': '1A2130', 'line': '28303F',
}
LAYOUTS = {'cover', 'evidence', 'comparison', 'process', 'timeline', 'closing'}
RICH_LAYOUTS = {'evidence_board', 'problem_solution', 'product_journey', 'market_model',
                'competition_matrix', 'milestone_roadmap', 'roadmap_funding', 'team_evidence'}
SECTION_LABELS = {
    'cover': '표지', 'background': '문제 배경 및 최근 트렌드', 'problem': '문제 정의',
    'alternatives': '기존 대체재의 문제점', 'solution': '해결방안', 'product': '제품 및 서비스 소개',
    'traction': '초기 고객 반응', 'business_model': '수익 모델', 'market': '시장 규모',
    'go_to_market': '초기 시장 진입 전략', 'growth': '성장 전략', 'milestones': '마일스톤',
    'team': '팀 구성', 'vision': '비전',
    'competition': '기존 대체재의 문제점',  # Legacy section alias.
}
RIGHTS = {'owned', 'licensed', 'permission', 'public-domain'}
KINDS = {'fact', 'plan', 'assumption'}
WIDTH, HEIGHT = 13 + 1 / 3, 7.5
MARGIN, CONTENT_W = .75, WIDTH - 1.5


def load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding='utf-8-sig'))
    except (OSError, ValueError) as exc:
        raise RenderError(f'{path.name}: JSON을 읽을 수 없습니다: {exc}') from exc
    if not isinstance(value, dict):
        raise RenderError(f'{path.name}: 최상위 값은 JSON object여야 합니다.')
    return value


def contained_path(project: Path, relative: str, *, must_exist: bool = True) -> Path:
    """Resolve symlinks and prevent all input/output paths from escaping project."""
    if not isinstance(relative, str) or not relative or '\x00' in relative:
        raise RenderError('비어 있거나 잘못된 파일 경로입니다.')
    if re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:', relative) or relative.startswith(('\\', '/')):
        raise RenderError(f'프로젝트 내부 상대경로만 허용합니다: {relative}')
    # Backslashes are separators even when an exported project is used on Unix.
    safe_relative = relative.replace('\\', '/')
    resolved = (project / safe_relative).resolve()
    if not resolved.is_relative_to(project.resolve()):
        raise RenderError(f'프로젝트 경계를 벗어난 경로입니다: {relative}')
    if must_exist and not resolved.is_file():
        raise RenderError(f'파일이 없습니다: {relative}')
    return resolved


def _font_directories() -> list[Path]:
    dirs = [Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts']
    local_app = os.environ.get('LOCALAPPDATA')
    if local_app:
        dirs.append(Path(local_app) / 'Microsoft/Windows/Fonts')
    dirs.extend([Path.home() / '.local/share/fonts', Path.home() / '.fonts',
                 Path('/usr/share/fonts'), Path('/usr/local/share/fonts'),
                 Path('/System/Library/Fonts'), Path('/Library/Fonts'),
                 Path.home() / 'Library/Fonts'])
    return [p for p in dirs if p.is_dir()]


def resolve_font(profile: dict[str, Any]) -> dict[str, Any]:
    typography = profile.get('typography', {})
    preferred = typography.get('preferred_family', 'Pretendard')
    wanted = [preferred] + typography.get('fallback_families', [
        'Noto Sans CJK KR', 'Noto Sans KR', 'Malgun Gothic', '맑은 고딕', 'Apple SD Gothic Neo'])
    found: list[tuple[str, Path, str]] = []
    for directory in _font_directories():
        for path in sorted(directory.rglob('*')):
            if path.suffix.lower() not in {'.ttf', '.otf', '.ttc'}:
                continue
            try:
                font = ImageFont.truetype(str(path), 24)
                family, style = font.getname()
                found.append((family, path, style))
            except (OSError, ValueError):
                continue
    for requested in wanted:
        matches = [x for x in found if x[0].casefold() == requested.casefold()
                   or x[0].casefold().startswith(requested.casefold() + ' ')]
        if matches:
            # Prefer ordinary text faces over bold/slim faces, retaining variable font support.
            matches.sort(key=lambda x: (x[2].casefold() not in {'regular', 'normal'}, len(x[0])))
            family, path, style = matches[0]
            faces = list(dict.fromkeys(str(item[1]) for item in matches
                                       if item[0].casefold() == family.casefold()))
            return {'requested': preferred, 'selected': family, 'path': str(path), 'paths': faces,
                    'face': style, 'substituted': not family.casefold().startswith(preferred.casefold()),
                    'embedded': False}
    raise RenderError('한글 지원 글꼴을 찾지 못했습니다. Pretendard, Noto Sans KR 또는 맑은 고딕을 설치한 뒤 다시 실행하세요.')


def measured_lines(text: str, font_path: str, size_pt: float, width_in: float) -> list[str]:
    """Greedy measured wrapping, preserving explicit newlines and all content."""
    font = ImageFont.truetype(font_path, round(size_pt * 96 / 72))
    # Use a regular instance for variable fonts to avoid measuring Thin then rendering Regular.
    try:
        for name in font.get_variation_names():
            if name.decode('utf-8', errors='ignore').casefold() == 'regular':
                font.set_variation_by_name(name)
                break
    except (AttributeError, OSError):
        pass
    max_width = width_in * 96 * .94  # includes bold and renderer variation headroom
    lines: list[str] = []
    for explicit in text.split('\n'):
        if not explicit:
            lines.append('')
            continue
        line = ''
        for ch in explicit:
            if line and font.getlength(line + ch) > max_width:
                # Move a partial Latin word as one piece whenever possible.
                if ch.isascii() and ch.isalnum() and ' ' in line:
                    prefix, tail = line.rsplit(' ', 1)
                    if tail and tail.isascii() and len(tail) < 20:
                        lines.append(prefix.rstrip())
                        line = tail + ch
                        continue
                lines.append(line.rstrip())
                line = ch.lstrip()
            else:
                line += ch
        lines.append(line.rstrip())
    return lines


def evidence_refs(slide: dict[str, Any]) -> list[str]:
    refs = list(slide.get('evidence_ids', []))
    refs += slide.get('governing_evidence_ids', [])
    for body in slide.get('body', []):
        refs += body.get('evidence_ids', [])
    refs += slide.get('visual', {}).get('evidence_ids', [])
    for node in walk_objects(slide.get('blocks', [])):
        refs += node.get('evidence_ids', [])
    return list(dict.fromkeys(refs))


def walk_objects(value: Any):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walk_objects(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk_objects(child)


def select_view(deck: dict[str, Any], view: str) -> dict[str, Any]:
    """Select a presentation without mutating the shared master/pitch slide pool."""
    if not deck.get('views'):
        return {**deck, 'view': view}
    spec = deck['views'].get(view)
    if not isinstance(spec, dict) or not isinstance(spec.get('slide_ids'), list) or not spec['slide_ids']:
        raise RenderError(f'view {view}: slide_ids가 필요합니다.')
    by_id = {s.get('id'): s for s in deck.get('slides', [])}
    ids = spec['slide_ids']
    if len(set(ids)) != len(ids) or any(sid not in by_id for sid in ids):
        raise RenderError(f'view {view}: 중복되거나 존재하지 않는 slide ID입니다.')
    return {**deck, 'slides': [by_id[sid] for sid in ids], 'view': view}


def validate_rich_assets(project: Path, slide: dict[str, Any]) -> list[str]:
    warnings = []
    for asset in walk_objects(slide.get('blocks', [])):
        if 'path' not in asset:
            continue
        if asset.get('rights') not in RIGHTS:
            raise RenderError(f'{slide.get("id", "")}: 이미지 사용권을 확인하세요.')
        for field in ('path', 'caption', 'source', 'alt'):
            if not isinstance(asset.get(field), str) or not asset[field].strip():
                raise RenderError(f'이미지 {field}가 필요합니다.')
        image_path = contained_path(project, asset['path'])
        crop = asset.get('crop')
        if crop is not None:
            if (not isinstance(crop, dict) or any(isinstance(crop.get(k), bool) or not isinstance(crop.get(k), (int, float))
                    or not math.isfinite(crop[k]) or not 0 <= crop[k] < 1 for k in ('left', 'top', 'right', 'bottom'))
                    or crop['left'] + crop['right'] >= 1 or crop['top'] + crop['bottom'] >= 1):
                raise RenderError('crop은 유효한 원본 잘라내기 비율이어야 합니다 (left/top/right/bottom).')
        try:
            with Image.open(image_path) as im:
                im.verify()
            with Image.open(image_path) as im:
                if im.format not in {'PNG', 'JPEG'}:
                    raise RenderError('이미지는 PNG/JPEG로 준비하세요.')
                if max(im.size) < 1200:
                    warnings.append(f'{slide.get("id", "")}: {asset["path"]} 이미지 긴 변 1200px 미만')
        except (UnidentifiedImageError, OSError) as exc:
            raise RenderError(f'이미지를 읽을 수 없습니다: {asset["path"]}') from exc
    return warnings


def validate_rich_slide(project: Path, slide: dict[str, Any], sources: dict[str, Any]) -> list[str]:
    if slide.get('layout') not in LAYOUTS | RICH_LAYOUTS:
        raise RenderError(f'{slide.get("id")}: 지원하지 않는 layout입니다.')
    if not isinstance(slide.get('governing_message'), str) or not slide['governing_message'].strip():
        raise RenderError('슬라이드 제목이 필요합니다.')
    if slide.get('governing_kind') not in KINDS:
        raise RenderError('제목 kind가 필요합니다.')
    if slide['governing_kind'] == 'fact' and not slide.get('governing_evidence_ids'):
        raise RenderError('사실 제목에 근거가 필요합니다.')
    for ref in evidence_refs(slide):
        if ref not in sources:
            raise RenderError(f'{slide.get("id")}: 존재하지 않는 evidence ID: {ref}')
    for ref in slide.get('governing_evidence_ids', []):
        if slide['governing_kind'] == 'fact' and (sources[ref].get('status') != 'confirmed' or sources[ref].get('kind') == 'assumption'):
            raise RenderError('사실 제목에 확인된 근거가 필요합니다.')
    seconds = slide.get('seconds')
    if isinstance(seconds, bool) or not isinstance(seconds, (int, float)) or not math.isfinite(seconds) or seconds < 0:
        raise RenderError('seconds는 0 이상의 유한 숫자여야 합니다.')
    if not slide.get('speaker_notes'):
        raise RenderError('발표 원고가 필요합니다.')
    quality_path = Path(__file__).with_name('quality.py')
    spec = importlib.util.spec_from_file_location('seed_ir_quality_renderer', quality_path)
    quality = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(quality)
    errors = quality.validate_slide_blocks(slide, sources)
    if errors:
        raise RenderError('\n'.join(errors))
    return validate_rich_assets(project, slide)


def disclosed_text(value: str, kind: str) -> str:
    """Keep an existing hypothesis/plan disclosure without adding a duplicate."""
    terms = {'plan': ('목표', '계획', '예정', '검증할'), 'assumption': ('가정', '추정', '예시')}
    if kind == 'fact' or any(term in value for term in terms.get(kind, ())):
        return value
    return {'plan': '계획 · ', 'assumption': '가정 · '}[kind] + value


def validate_contract(project: Path, deck: dict[str, Any], evidence: dict[str, Any],
                      profile: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    slides = deck.get('slides')
    if not isinstance(slides, list) or not slides:
        raise RenderError('deck/deck.json의 slides에 최소 1개의 슬라이드가 필요합니다.')
    if not isinstance(deck.get('team_name'), str) or not deck['team_name'].strip():
        raise RenderError('deck/deck.json에 team_name이 필요합니다.')
    items = evidence.get('items', [])
    if not isinstance(items, list):
        raise RenderError('evidence/evidence.json의 items는 배열이어야 합니다.')
    sources = {x['id']: x for x in items if isinstance(x, dict) and isinstance(x.get('id'), str)}
    if len(sources) != len(items):
        raise RenderError('evidence ID가 없거나 중복되었습니다.')
    warnings: list[str] = []
    seen = set()
    for number, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            raise RenderError(f'{number}페이지: slide는 JSON object여야 합니다.')
        label = slide.get('id', f'slide-{number}')
        if not isinstance(label, str) or not label.strip():
            raise RenderError(f'{number}페이지: slide ID는 비어 있지 않은 문자열이어야 합니다.')
        if label in seen:
            raise RenderError(f'중복 slide ID: {label}')
        seen.add(label)
        if 'blocks' in slide:
            warnings.extend(validate_rich_slide(project, slide, sources))
            continue
        if slide.get('layout') not in LAYOUTS:
            raise RenderError(f'{label}: layout은 {sorted(LAYOUTS)} 중 하나여야 합니다.')
        message = slide.get('governing_message', '')
        if not isinstance(message, str) or not message.strip() or len(message) > 66:
            raise RenderError(f'{label}: governing_message는 공백이 아닌 1–66자여야 합니다. 핵심 주장 1문장으로 줄이세요.')
        body = slide.get('body', [])
        if not isinstance(body, list) or len(body) > 3:
            raise RenderError(f'{label}: 본문은 최대 3개입니다. 부록 또는 다음 슬라이드로 나누세요.')
        if slide['layout'] == 'cover' and len(body) > 2:
            raise RenderError(f'{label}: 표지 본문은 최대 2개입니다.')
        for i, block in enumerate([{'text': message, 'kind': slide.get('governing_kind'),
                                     'evidence_ids': slide.get('governing_evidence_ids', [])}] + body):
            if not isinstance(block, dict) or not isinstance(block.get('text'), str):
                raise RenderError(f'{label}: 본문 {i}은 text, kind, evidence_ids가 있는 object여야 합니다.')
            if block.get('kind') not in KINDS:
                raise RenderError(f'{label}: 주장 kind는 fact, plan, assumption 중 하나여야 합니다.')
            if i and (not block['text'].strip() or len(block['text']) > 80):
                raise RenderError(f'{label}: 본문 {i}은 1–80자여야 합니다. 내용 축약 또는 슬라이드 분리가 필요합니다.')
            if not isinstance(block.get('evidence_ids', []), list):
                raise RenderError(f'{label}: evidence_ids는 배열이어야 합니다.')
            if any(not isinstance(ref, str) or not ref for ref in block.get('evidence_ids', [])):
                raise RenderError(f'{label}: evidence_ids는 비어 있지 않은 문자열의 배열이어야 합니다.')
            if block['kind'] == 'fact' and not block.get('evidence_ids'):
                raise RenderError(f'{label}: 사실 주장에 evidence_ids가 없습니다: {block["text"]}')
            if block['kind'] == 'fact':
                for ref in block.get('evidence_ids', []):
                    source = sources.get(ref)
                    if source and (source.get('status') != 'confirmed' or source.get('kind') == 'assumption'):
                        raise RenderError(f'{label}: 사실 주장 [{ref}]에는 확인된 내부/외부 근거가 필요합니다. 가정·미검증 근거를 사실로 내보낼 수 없습니다.')
        for field in ('evidence_ids', 'governing_evidence_ids'):
            if not isinstance(slide.get(field, []), list) or any(not isinstance(x, str) or not x for x in slide.get(field, [])):
                raise RenderError(f'{label}: {field}는 비어 있지 않은 문자열의 배열이어야 합니다.')
        raw_visual = slide.get('visual', {})
        if not isinstance(raw_visual, dict) or not isinstance(raw_visual.get('evidence_ids', []), list) or any(not isinstance(x, str) or not x for x in raw_visual.get('evidence_ids', [])):
            raise RenderError(f'{label}: visual.evidence_ids는 비어 있지 않은 문자열의 배열이어야 합니다.')
        for ref in evidence_refs(slide):
            if ref not in sources:
                raise RenderError(f'{label}: 존재하지 않는 evidence ID: {ref}')
        seconds = slide.get('seconds')
        if (isinstance(seconds, bool) or not isinstance(seconds, (int, float)) or not math.isfinite(seconds)
                or seconds < 0 or (seconds == 0 and not slide.get('appendix'))):
            raise RenderError(f'{label}: seconds는 양수여야 합니다. appendix만 0초를 허용합니다.')
        assets = slide.get('assets', [])
        if not isinstance(assets, list):
            raise RenderError(f'{label}: assets는 배열이어야 합니다.')
        limit = {'cover': 1, 'evidence': 1, 'comparison': 0, 'process': 3, 'timeline': 0, 'closing': 1}[slide['layout']]
        if len(assets) > limit:
            raise RenderError(f'{label}: {slide["layout"]} 레이아웃은 이미지 최대 {limit}개입니다. 다른 레이아웃으로 옮기세요.')
        if slide['layout'] == 'process' and len(assets) > len(body):
            raise RenderError(f'{label}: 과정 이미지는 해당하는 본문 단계가 있어야 합니다.')
        for a in assets:
            if not isinstance(a, dict) or a.get('rights') not in RIGHTS:
                raise RenderError(f'{label}: 이미지 사용권은 owned, licensed, permission, public-domain 중 확인된 값이어야 합니다. unknown 이미지는 내보내지 않습니다.')
            for required in ('path', 'caption', 'source', 'alt'):
                if not isinstance(a.get(required), str) or not a[required].strip():
                    raise RenderError(f'{label}: 이미지의 {required}를 기록하세요.')
            if len(a['caption']) > 80:
                raise RenderError(f'{label}: 이미지 캡션은 80자 이내로 줄이세요.')
            path = contained_path(project, a['path'])
            try:
                with Image.open(path) as image:
                    image.verify()
                with Image.open(path) as image:
                    if image.format not in {'PNG', 'JPEG'}:
                        raise RenderError(f'{label}: 이미지는 PNG/JPEG로 준비하세요: {a["path"]}')
                    if max(image.size) < 1200:
                        warnings.append(f'{label}: 이미지 {a["path"]}의 긴 변이 1200px 미만입니다. 실제 렌더에서 선명도를 확인하세요.')
            except (UnidentifiedImageError, OSError) as exc:
                raise RenderError(f'{label}: 이미지를 읽을 수 없습니다: {a["path"]}') from exc
        visual = slide.get('visual', {'type': 'none'})
        if not isinstance(visual, dict) or visual.get('type', 'none') not in {'bar', 'metrics', 'none'}:
            raise RenderError(f'{label}: visual.type은 bar, metrics, none 중 하나여야 합니다.')
        if visual.get('type') in {'bar', 'metrics'}:
            categories, values = visual.get('categories', []), visual.get('values', [])
            if (not isinstance(categories, list) or not isinstance(values, list) or not categories
                    or len(categories) != len(values) or len(values) > (3 if visual['type'] == 'metrics' else 6)):
                raise RenderError(f'{label}: 차트 범주와 값의 수가 같아야 합니다. metrics 최대 3개, bar 최대 6개입니다.')
            if any(not isinstance(x, str) or not x.strip() or len(x) > 20 for x in categories):
                raise RenderError(f'{label}: 차트 범주명은 1–20자 문자열이어야 합니다.')
            if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in values):
                raise RenderError(f'{label}: 차트 values는 유한 숫자여야 합니다.')
            if not visual.get('source') or not isinstance(visual.get('unit'), str) or not visual.get('unit').strip() or not visual.get('evidence_ids'):
                raise RenderError(f'{label}: 차트에 source, unit, evidence_ids를 기록하세요. 목표/가정이면 출처 항목도 해당 kind로 기록하세요.')
            if slide['layout'] not in {'evidence', 'closing'}:
                raise RenderError(f'{label}: 차트는 evidence 또는 closing 레이아웃을 사용하세요.')
            if assets:
                raise RenderError(f'{label}: 차트와 이미지는 한 슬라이드의 같은 시각 영역을 공유할 수 없습니다. 분리하세요.')
    total = sum(s['seconds'] for s in slides)
    if deck.get('view', 'pitch') == 'pitch' and not 270 <= total <= 300:
        warnings.append(f'발화 계획은 {total:g}초입니다. 5분 최종본의 목표는 270–300초이며 실제 리허설을 따로 측정해야 합니다.')
    return sources, warnings


class DeckRenderer:
    def __init__(self, project: Path, deck: dict[str, Any], profile: dict[str, Any],
                 sources: dict[str, Any], font: dict[str, Any]):
        self.project, self.deck, self.profile, self.sources, self.font = project, deck, profile, sources, font
        self.colors = {**DEFAULT_COLORS, **profile.get('colors', {})}
        if any(not isinstance(v, str) or not re.fullmatch(r'#?[0-9a-fA-F]{6}', v) for v in self.colors.values()):
            raise RenderError('design-profile.json 색상은 6자리 RGB hex여야 합니다.')
        self.colors = {k: v.lstrip('#') for k, v in self.colors.items()}
        typ = profile.get('typography', {})
        self.sizes = {key: float(typ.get(key, default)) for key, default in {
            'title_pt': 32, 'body_pt': 20, 'metric_pt': 42, 'caption_pt': 12, 'footer_pt': 9,
        }.items()}
        if any(not math.isfinite(v) or v < 8 or v > 60 for v in self.sizes.values()):
            raise RenderError('글꼴 크기는 8–60pt 사이의 유한 숫자여야 합니다.')
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
        self.prs.core_properties.title = f'{deck["team_name"]} Seed IR'
        self.prs.core_properties.subject = 'Evidence-linked Seed investor pitch draft'
        self.prs.core_properties.author = deck['team_name']
        self.prs.core_properties.keywords = 'Seed IR, draft, evidence-linked'
        self.layout_checks: list[dict[str, Any]] = []

    def color(self, key: str) -> RGBColor:
        return RGBColor.from_string(self.colors.get(key, key).lstrip('#'))

    def text(self, slide: Any, value: str, x: float, y: float, w: float, h: float,
             *, size: float = 20, color: str = 'ink', bold: bool = False,
             max_lines: int | None = None, role: str = 'body', slide_id: str = '',
             highlight_numbers: bool = False) -> Any:
        lines = measured_lines(value, self.font['path'], size, w - .04)
        line_height = size * 1.28 / 72
        allowed = max(1, math.floor((h - .025) / line_height))
        if max_lines is not None:
            allowed = min(allowed, max_lines)
        if len(lines) > allowed:
            raise RenderError(f'{slide_id}: {role}가 {len(lines)}줄로 영역 {allowed}줄을 초과합니다. 폰트 축소 대신 내용을 줄이거나 슬라이드를 분리하세요: {value}')
        if x < -.001 or y < -.001 or x + w > WIDTH + .001 or y + h > HEIGHT + .001:
            raise RenderError(f'{slide_id}: {role} 영역이 슬라이드 경계를 벗어났습니다.')
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = Inches(.015)
        frame.margin_top = frame.margin_bottom = 0
        # Explicit breaks make our measured wrapping stable across office suites.
        for line_index, line in enumerate(lines):
            para = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
            para.space_after = Pt(0)
            para.space_before = Pt(0)
            para.line_spacing = 1.15
            parts = re.split(r'(\d[\d,.]*(?:%|억|만|개월|초|원)?)', line) if highlight_numbers else [line]
            for part in parts:
                run = para.add_run()
                run.text = part
                run.font.name, run.font.size, run.font.bold = self.font['selected'], Pt(size), bold
                run.font.color.rgb = self.color('accent' if highlight_numbers and re.match(r'^\d', part) else color)
                rpr = run._r.get_or_add_rPr()
                for tag in ('a:ea', 'a:cs'):
                    item = OxmlElement(tag)
                    item.set('typeface', self.font['selected'])
                    rpr.append(item)
        self.layout_checks.append({'slide': slide_id, 'role': role, 'lines': len(lines), 'allowed_lines': allowed,
                                   'font_pt': size, 'box_inches': [x, y, w, h]})
        return shape

    def line(self, slide: Any, x: float, y: float, w: float, color: str = 'accent') -> None:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(.018))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.color(color)
        shape.line.fill.background()

    def body_text(self, slide: Any, block: dict[str, Any], x: float, y: float, w: float, h: float,
                  *, dark: bool, sid: str, size: float = 20) -> None:
        size = self.sizes['body_pt'] + (size - 20)
        self.text(slide, disclosed_text(block['text'], block['kind']), x, y, w, h, size=size,
                  color='white' if dark else 'ink', slide_id=sid)

    def picture(self, slide: Any, asset: dict[str, Any], x: float, y: float, w: float, h: float,
                *, dark: bool, sid: str) -> None:
        path = contained_path(self.project, asset['path'])
        with Image.open(path) as im:
            iw, ih = im.size
        caption_h = .55
        scale = min(w / iw, (h - caption_h) / ih)
        pw, ph = iw * scale, ih * scale
        picture = slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y),
                                           width=Inches(pw), height=Inches(ph))
        picture._element.nvPicPr.cNvPr.set('descr', asset['alt'])
        self.text(slide, asset['caption'], x, y + h - caption_h, w, caption_h,
                  size=self.sizes['caption_pt'], color='muted_dark' if dark else 'muted', role='image caption', slide_id=sid)

    def visual(self, slide: Any, visual: dict[str, Any], x: float, y: float, w: float, h: float,
               *, dark: bool, sid: str) -> None:
        text_color = 'white' if dark else 'ink'
        if visual['type'] == 'metrics':
            count = len(visual['values'])
            item_h = h / count
            for i, (category, value) in enumerate(zip(visual['categories'], visual['values'])):
                self.text(slide, f'{value:,.2f}'.rstrip('0').rstrip('.') if isinstance(value, float) else f'{value:,}',
                          x, y + i * item_h, w * .65, min(.88, item_h * .63), size=self.sizes['metric_pt'],
                          color='accent', bold=True, slide_id=sid, role='metric')
                self.text(slide, f'{category} · {visual["unit"]}', x, y + i * item_h + min(.83, item_h * .61),
                          w, min(.5, item_h * .37), size=14, color=text_color, role='metric label', slide_id=sid)
        else:
            self.text(slide, '단위: ' + visual['unit'], x, y - .26, w, .25,
                      size=10, color='muted_dark' if dark else 'muted', role='chart unit', slide_id=sid)
            data = CategoryChartData()
            data.categories = visual['categories']
            data.add_series(visual['unit'], visual['values'])
            chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,
                    Inches(x), Inches(y), Inches(w), Inches(h), data).chart
            # python-pptx's historical templates use negative axis IDs. The OOXML
            # schema requires unsignedInt; normalize both identifiers and links
            # so strict importers can render the same editable chart.
            for node in chart._chartSpace.xpath('.//c:axId | .//c:crossAx'):
                node.set('val', str(int(node.get('val')) & 0xffffffff))
            chart.has_legend = False
            chart.has_title = False
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.font.name = self.font['selected']
            plot.data_labels.font.size = Pt(14)
            plot.data_labels.font.color.rgb = self.color(text_color)
            plot.data_labels.number_format = '#,##0' if all(float(v).is_integer() for v in visual['values']) else '#,##0.00'
            plot.gap_width = 65
            for axis in (chart.category_axis, chart.value_axis):
                axis.tick_labels.font.name = self.font['selected']
                axis.tick_labels.font.size = Pt(13)
                axis.tick_labels.font.color.rgb = self.color(text_color)
                axis.format.line.fill.background()
            chart.value_axis.has_major_gridlines = True
            chart.value_axis.major_gridlines.format.line.color.rgb = self.color('line' if dark else 'muted_dark')
            if min(visual['values']) >= 0:
                chart.value_axis.minimum_scale = 0
            for point in chart.series[0].points:
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self.color('accent')
                point.format.line.fill.background()
            # Transparent native chart background retains the reference paper/ink canvas.
            chart_space = chart._chartSpace
            sp_pr = chart_space.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr')
            if sp_pr is None:
                sp_pr = OxmlElement('c:spPr')
                chart_space.append(sp_pr)
            sp_pr.append(OxmlElement('a:noFill'))

    def build(self) -> Any:
        total = len(self.deck['slides'])
        for number, content in enumerate(self.deck['slides'], 1):
            sid = content.get('id', f'slide-{number}')
            layout, body, assets = content['layout'], content.get('body', []), content.get('assets', [])
            visual = content.get('visual', {'type': 'none'})
            dark = layout in {'cover', 'comparison', 'timeline', 'closing'}
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = self.color('ink' if dark else 'paper')
            section = str(content.get('section', 'Seed IR'))
            self.text(slide, SECTION_LABELS.get(section, section), MARGIN, .40, 10, .25,
                      size=11, color='gold' if dark else 'accent', bold=True, role='section', slide_id=sid)
            self.text(slide, disclosed_text(content['governing_message'], content['governing_kind']), MARGIN, .88, CONTENT_W, 1.28,
                      size=self.sizes['title_pt'], color='white' if dark else 'ink', bold=True,
                      max_lines=2, role='governing message', slide_id=sid, highlight_numbers=True)
            top, height = 2.48, 3.93
            has_visual = visual.get('type') in {'bar', 'metrics'}
            if layout in {'cover', 'closing'}:
                has_right = bool(assets or has_visual)
                body_w = 6.5 if has_right else 10.8
                if body:
                    row_h = min(1.23, height / len(body))
                    for i, block in enumerate(body):
                        self.body_text(slide, block, MARGIN, top + .15 + i * row_h, body_w,
                                       row_h - .08, dark=dark, sid=sid, size=22 if layout == 'cover' else 20)
                if assets:
                    self.picture(slide, assets[0], 8.05, top, 4.53, height, dark=dark, sid=sid)
                if has_visual:
                    self.visual(slide, visual, 8.05, top, 4.53, height, dark=dark, sid=sid)
                self.line(slide, MARGIN, 6.60, CONTENT_W, 'gold')
            elif layout == 'evidence':
                if has_visual or assets:
                    if has_visual:
                        self.visual(slide, visual, MARGIN, top, 6.55, height, dark=dark, sid=sid)
                    else:
                        self.picture(slide, assets[0], MARGIN, top, 6.55, height, dark=dark, sid=sid)
                    for i, block in enumerate(body):
                        row_h = height / max(len(body), 1)
                        self.body_text(slide, block, 7.75, top + i * row_h, 4.83, row_h - .14,
                                       dark=dark, sid=sid, size=19)
                else:
                    for i, block in enumerate(body):
                        row_h = height / max(len(body), 1)
                        self.text(slide, f'{i + 1:02}', MARGIN, top + i * row_h, .65, .60,
                                  size=27, color='accent', bold=True, role='finding number', slide_id=sid)
                        self.body_text(slide, block, 1.65, top + .08 + i * row_h, 10.92,
                                       row_h - .15, dark=dark, sid=sid, size=22)
                        self.line(slide, 1.65, top + (i + 1) * row_h - .15, 10.92, 'muted_dark')
            elif layout in {'comparison', 'process', 'timeline'}:
                count = max(len(body), 1)
                gap = .40
                col_w = (CONTENT_W - gap * (count - 1)) / count
                if layout == 'timeline':
                    self.line(slide, MARGIN, top + .60, CONTENT_W, 'gold')
                for i, block in enumerate(body):
                    x = MARGIN + i * (col_w + gap)
                    self.text(slide, f'{i + 1:02}', x, top, col_w, .53, size=25,
                              bold=True, color='gold' if dark else 'accent', role='stage number', slide_id=sid)
                    if layout == 'comparison':
                        self.line(slide, x, top + .61, col_w, 'accent' if i == count - 1 else 'gold')
                    has_asset = layout == 'process' and i < len(assets)
                    text_h = 1.46 if has_asset else height - .90
                    self.body_text(slide, block, x, top + .85, col_w, text_h,
                                   dark=dark, sid=sid, size=20)
                    if has_asset:
                        self.picture(slide, assets[i], x, top + 2.37, col_w, 1.60,
                                     dark=dark, sid=sid)
            refs = evidence_refs(content)
            footer_parts = []
            for ref in refs:
                source = self.sources[ref]
                label = str(source.get('publisher') or {
                    'internal': '팀 제출자료', 'external': '외부 자료', 'assumption': '가정·추정',
                }.get(source.get('kind'), '근거 자료'))
                if source.get('locator'):
                    label += ' · ' + str(source['locator'])
                footer_parts.append(f'[{ref}] {label[:35] + ("…" if len(label) > 35 else "")}')
            if assets:
                footer_parts.append('이미지 출처·사용권: 발표자 노트')
            if not refs:
                footer_parts.append('계획·가정: 발표자 노트 확인')
            self.text(slide, '  '.join(footer_parts), MARGIN, 6.98, 10.45, .35,
                      size=self.sizes['footer_pt'], color='muted_dark' if dark else 'muted', max_lines=2,
                      role='source footer', slide_id=sid)
            self.text(slide, f'{number:02} / {total:02}', 11.67, 7.02, .91, .23,
                      size=9, color='muted_dark' if dark else 'muted', role='page number', slide_id=sid)
            notes = [f'Slide ID: {sid}', f'계획 발화 시간: {content["seconds"]}초',
                     f'제목 상태: {content["governing_kind"]}', '', '발표 원고',
                     str(content.get('speaker_notes', '')), '', '근거 추적']
            for ref in refs:
                notes.append(f'[{ref}] ' + json.dumps(self.sources[ref], ensure_ascii=False))
            if has_visual:
                notes += ['', '차트 데이터·단위·출처', json.dumps(visual, ensure_ascii=False)]
            for i, asset in enumerate(assets, 1):
                notes += ['', f'이미지 {i}', json.dumps(asset, ensure_ascii=False)]
            notes += ['', '내보내기 상태: 초안. 실제 PPTX 렌더와 5분 리허설을 별도로 확인해야 합니다.']
            slide.notes_slide.notes_text_frame.text = '\n'.join(notes)
        return self.prs


def write_storyboard(project: Path, deck: dict[str, Any], sources: dict[str, Any], path: Path) -> None:
    parts = ['<!doctype html><html lang="ko"><meta charset="utf-8">',
             '<meta name="viewport" content="width=device-width,initial-scale=1">',
             '<title>Seed IR 내용 검토</title><style>',
             'body{margin:0;background:#e9e5dd;color:#10141b;font-family:"Pretendard","Noto Sans KR","Malgun Gothic",sans-serif}header{padding:32px;max-width:1200px;margin:auto}h1{margin:0 0 12px}header p{line-height:1.7}main{max-width:1200px;margin:auto;padding:0 24px 48px;display:grid;gap:24px}article{padding:36px;background:#f5f2eb;border-top:4px solid #e0492e}article.dark{background:#10141b;color:#fff}small{color:#c57f1e}h2{font-size:30px;line-height:1.35;max-width:1000px}li{font-size:20px;line-height:1.6;margin:12px 0}.media{display:flex;gap:20px;flex-wrap:wrap}figure{margin:12px 0;max-width:31%}img{max-width:100%;max-height:300px;object-fit:contain}figcaption{font-size:13px;line-height:1.5}details{margin-top:24px;border-top:1px solid #87909b;padding-top:16px}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:14px/1.7 inherit}table{border-collapse:collapse}th,td{padding:12px 24px;border-bottom:1px solid #87909b;text-align:left}.tag{font-size:13px;color:#c57f1e}a{color:inherit}@media(max-width:650px){article{padding:22px}h2{font-size:24px}figure{max-width:100%}}@media print{article{break-after:page}details{display:block}}',
             '</style><header><h1>' + html.escape(deck['team_name']) + ' · Seed IR 내용 검토</h1>',
             '<p>이 HTML은 검토용 storyboard입니다. PPTX를 렌더한 이미지가 아닙니다. 슬라이드의 실제 줄바꿈·겹침·글꼴과 이미지 선명도는 PowerPoint 또는 사용 중인 발표 앱에서 확인하세요.</p>',
             f'<p>{len(deck["slides"])}장 · 발화 계획 {sum(s["seconds"] for s in deck["slides"]):g}초</p></header><main>']
    for i, slide in enumerate(deck['slides'], 1):
        dark = slide['layout'] in {'cover', 'closing', 'comparison', 'timeline'}
        section = str(slide.get('section', ''))
        parts += [f'<article class="{"dark" if dark else "light"}"><small>{i:02} · {html.escape(SECTION_LABELS.get(section, section))} · {slide["seconds"]}초</small>',
                  f'<h2>{html.escape(slide["governing_message"])}</h2><div class="tag">{html.escape(slide["governing_kind"])}</div><ul>']
        for body in slide.get('body', []):
            parts.append('<li>' + html.escape(disclosed_text(body['text'], body['kind'])) + '</li>')
        for block in slide.get('blocks', []):
            parts.append('<li><strong>' + html.escape(disclosed_text(block.get('title', ''), block.get('kind', 'fact'))) + '</strong><pre>' + html.escape(json.dumps(block, ensure_ascii=False, indent=2)) + '</pre></li>')
        parts.append('</ul>')
        visual = slide.get('visual', {})
        if visual.get('type') in {'bar', 'metrics'}:
            parts.append('<table><thead><tr><th>범주</th><th>값 (' + html.escape(visual['unit']) + ')</th></tr></thead><tbody>')
            for category, value in zip(visual['categories'], visual['values']):
                parts.append(f'<tr><td>{html.escape(category)}</td><td>{value:,}</td></tr>')
            parts.append('</tbody></table>')
        parts.append('<div class="media">')
        for asset in slide.get('assets', []):
            source = contained_path(project, asset['path'])
            with Image.open(source) as im:
                mime = 'image/png' if im.format == 'PNG' else 'image/jpeg'
            uri = 'data:' + mime + ';base64,' + base64.b64encode(source.read_bytes()).decode('ascii')
            parts.append(f'<figure><img src="{uri}" alt="{html.escape(asset["alt"], quote=True)}"><figcaption>{html.escape(asset["caption"])}</figcaption></figure>')
        parts.append('</div><details><summary>발표 원고와 근거</summary><pre>')
        detail = [str(slide.get('speaker_notes', '')), '', '근거']
        for ref in evidence_refs(slide):
            detail.append('[' + ref + '] ' + json.dumps(sources[ref], ensure_ascii=False))
        if slide.get('assets'):
            detail += ['', '이미지 사용권', json.dumps(slide['assets'], ensure_ascii=False, indent=2)]
        parts += [html.escape('\n'.join(detail)), '</pre></details></article>']
    parts.append('</main></html>')
    path.write_text(''.join(parts), encoding='utf-8')


def artifact_runtime() -> tuple[Path, Path]:
    """Discover Codex's bundled runtime without embedding a machine/user path."""
    candidates: list[Path] = []
    if os.environ.get('SEED_IR_NODE'):
        candidates.append(Path(os.environ['SEED_IR_NODE']))
    found = shutil.which('node')
    if found:
        candidates.append(Path(found))
    cache = Path.home() / '.cache/codex-runtimes'
    if cache.is_dir():
        for directory in sorted(cache.iterdir()):
            candidates += [directory / 'dependencies/node/bin/node.exe', directory / 'dependencies/node/bin/node']
    modules_override = os.environ.get('SEED_IR_NODE_MODULES')
    for node in candidates:
        if not node.is_file():
            continue
        package_roots = ([Path(modules_override)] if modules_override else []) + [
            node.parent.parent / 'node_modules', node.parent / 'node_modules',
            Path.cwd() / 'node_modules']
        for root in package_roots:
            module = root / '@oai/artifact-tool/dist/artifact_tool.mjs'
            if module.is_file():
                return node, module
    raise RenderError('고급 장표에는 Codex의 @oai/artifact-tool 런타임이 필요합니다. Codex workspace dependencies를 확인하거나 SEED_IR_NODE와 SEED_IR_NODE_MODULES를 설정하세요. 내용이 줄어드는 간이 렌더로 대체하지 않았습니다.')


def chart_snapshot_helper() -> Path:
    helper = Path(__file__).with_name('materialize_chart_workbooks.py')
    if not helper.is_file():
        raise RenderError('배포 파일 materialize_chart_workbooks.py가 없습니다. 플러그인을 다시 설치하세요.')
    return helper


def render_rich_project(project: Path, deck: dict[str, Any], profile: dict[str, Any],
                        sources: dict[str, Any], warnings: list[str], font: dict[str, Any]) -> dict[str, Any]:
    node, module = artifact_runtime()
    view = deck.get('view', 'pitch')
    output = contained_path(project, f'output/{view}/seed-ir-draft.pptx', must_exist=False).parent
    output.mkdir(parents=True, exist_ok=True)
    # Complete generation and reimport rendering before replacing prior deliverables.
    # Keep intermediate packages outside a synced project so its watcher does
    # not lock the export/repair files. Cleanup failure is not render failure.
    with tempfile.TemporaryDirectory(prefix='seed-ir-rich-', ignore_cleanup_errors=True) as tmp:
        temp = Path(tmp)
        request_path = temp / 'render-input.json'
        has_charts = any(b.get('type') == 'chart' for s in deck['slides'] for b in s.get('blocks', [])) or any(s.get('visual', {}).get('type') == 'bar' for s in deck['slides'])
        request_path.write_text(json.dumps({'deck': deck, 'profile': profile, 'sources': sources,
            'project': str(project), 'font': font, 'section_labels': SECTION_LABELS,
            'chart_snapshot_helper': str(chart_snapshot_helper()) if has_charts else None}, ensure_ascii=False), encoding='utf-8')
        result = subprocess.run([str(node), str(Path(__file__).with_name('render_rich.mjs')),
                                 str(request_path), str(temp), str(module), sys.executable], capture_output=True, text=True,
                                encoding='utf-8', errors='replace')
        if result.returncode:
            raise RenderError('고급 장표 렌더 실패:\n' + result.stderr[-9000:])
        if result.stdout:
            print(result.stdout, end='', file=sys.stderr)
        rich_report = load_json(temp / 'rich-render-report.json')
        # Reopen native objects independently of the authoring library.
        actual = Presentation(str(temp / 'seed-ir-draft.pptx'))
        actual_charts = sum(shape.has_chart for s in actual.slides for shape in s.shapes)
        actual_tables = sum(shape.has_table for s in actual.slides for shape in s.shapes)
        expected_charts = sum(b.get('type') == 'chart' for s in deck['slides'] for b in s.get('blocks', []))
        expected_tables = sum(b.get('type') == 'table' for s in deck['slides'] for b in s.get('blocks', []))
        if actual_charts < expected_charts or actual_tables < expected_tables:
            raise RenderError('편집 가능한 차트/표 수가 계약보다 적습니다.')
        write_storyboard(project, deck, sources, temp / 'storyboard.html')
        for filename in ('seed-ir-draft.pptx', 'storyboard.html', 'rich-render-report.json'):
            shutil.copy2(temp / filename, output / filename)
        if (temp / 'chart-data-snapshot.json').is_file():
            private_report = contained_path(project, f'work/render/{view}-chart-data-snapshot.json', must_exist=False)
            private_report.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(temp / 'chart-data-snapshot.json', private_report)
        if (temp / 'media-dedup-report.json').is_file():
            private_report = contained_path(project, f'work/render/{view}-media-dedup-report.json', must_exist=False)
            private_report.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(temp / 'media-dedup-report.json', private_report)
        renders = output / 'renders'
        renders.mkdir(exist_ok=True)
        for source in (temp / 'renders').iterdir():
            shutil.copy2(source, renders / source.name)
    report = {'status': 'draft_exported', 'engine': 'artifact-tool', 'view': view,
              'source_deck_sha256': deck.get('source_deck_sha256'),
              'slide_count': len(deck['slides']), 'planned_seconds': sum(s.get('seconds', 0) for s in deck['slides']),
              'font': {k: v for k, v in font.items() if k not in {'path', 'paths'}}, 'warnings': warnings,
              'editable_chart_count': actual_charts, 'editable_table_count': actual_tables,
              'chart_data_packaging': 'literal_input_workbook_snapshot' if actual_charts else 'not_applicable',
              'pptx_render_review': 'not_performed', 'actual_pptx_rendered': True,
              'storyboard_is_pptx_render': False, 'layout_preflight': rich_report['checks'],
              'block_coverage': rich_report['coverage'],
              'compositions': rich_report.get('compositions', []),
              'outputs': [f'output/{view}/seed-ir-draft.pptx', f'output/{view}/storyboard.html'],
              'next_step': '실제 렌더의 정보 위계·근거 충실도·레이아웃 다양성을 검수하세요.'}
    (output / 'render-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    if view == 'pitch':
        for filename in ('seed-ir-draft.pptx', 'storyboard.html', 'render-report.json'):
            shutil.copy2(output / filename, output.parent / filename)
    return report


def render_project(project_dir: str | Path, view: str = 'pitch') -> dict[str, Any]:
    project = Path(project_dir).expanduser().resolve()
    if not project.is_dir():
        raise RenderError(f'프로젝트 폴더가 없습니다: {project}')
    source_bytes = contained_path(project, 'deck/deck.json').read_bytes()
    source_deck = json.loads(source_bytes.decode('utf-8-sig'))
    if not isinstance(source_deck, dict):
        raise RenderError('deck.json 최상위는 object여야 합니다.')
    deck = select_view(source_deck, view)
    deck['source_deck_sha256'] = hashlib.sha256(source_bytes).hexdigest()
    profile = load_json(contained_path(project, 'design/design-profile.json'))
    evidence = load_json(contained_path(project, 'evidence/evidence.json'))
    sources, warnings = validate_contract(project, deck, evidence, profile)
    font = resolve_font(profile)
    if font['substituted']:
        warnings.append(f'{font["requested"]} 미설치: {font["selected"]}로 대체했습니다. 실제 발표 컴퓨터의 글꼴과 줄바꿈을 확인하세요.')
    if deck.get('views') or any('blocks' in s for s in deck['slides']):
        return render_rich_project(project, deck, profile, sources, warnings, font)
    renderer = DeckRenderer(project, deck, profile, sources, font)
    prs = renderer.build()  # preflight all text geometry before writing any deliverable
    output = contained_path(project, 'output/seed-ir-draft.pptx', must_exist=False).parent
    report_path = contained_path(project, 'output/render-report.json', must_exist=False)
    output.mkdir(parents=True, exist_ok=True)
    pptx_path, story_path = output / 'seed-ir-draft.pptx', output / 'storyboard.html'
    # Same-directory temporary files support atomic replacement and preserve prior valid output.
    with tempfile.TemporaryDirectory(prefix='.seed-ir-', dir=output) as temp:
        temp_dir = Path(temp)
        candidate = temp_dir / 'seed-ir-draft.pptx'
        prs.save(str(candidate))
        # Verify editable slide count and native chart counts by re-opening the written file.
        reopened = Presentation(str(candidate))
        if len(reopened.slides) != len(deck['slides']):
            raise RenderError('생성된 PPTX의 슬라이드 수가 입력과 다릅니다.')
        expected_charts = sum(s.get('visual', {}).get('type') == 'bar' for s in deck['slides'])
        actual_charts = sum(shape.has_chart for s in reopened.slides for shape in s.shapes)
        if expected_charts != actual_charts:
            raise RenderError('편집 가능한 차트 수가 입력과 다릅니다.')
        write_storyboard(project, deck, sources, temp_dir / 'storyboard.html')
        candidate.replace(pptx_path)
        (temp_dir / 'storyboard.html').replace(story_path)
    report = {'status': 'draft_exported', 'slide_count': len(deck['slides']),
              'source_deck_sha256': deck.get('source_deck_sha256'),
              'planned_seconds': sum(s['seconds'] for s in deck['slides']),
              'font': {k: v for k, v in font.items() if k not in {'path', 'paths'}}, 'warnings': warnings,
              'editable_chart_count': actual_charts, 'pptx_render_review': 'not_performed',
              'storyboard_is_pptx_render': False, 'layout_preflight': renderer.layout_checks,
              'outputs': ['output/seed-ir-draft.pptx', 'output/storyboard.html'],
              'next_step': '실제 PPTX 렌더를 검수하고 reviews/visual.json에 검토 결과를 기록하세요.'}
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description='근거가 연결된 Seed IR JSON으로 편집 가능한 PPTX 초안을 만듭니다.')
    parser.add_argument('--project', type=Path)
    parser.add_argument('--view', '--variant', dest='view', choices=['master', 'pitch'], default='pitch')
    parser.add_argument('--check-runtime', action='store_true', help='고급 장표용 Node/Artifact Tool 런타임을 확인합니다.')
    args = parser.parse_args()
    try:
        if args.check_runtime:
            node, module = artifact_runtime()
            print(json.dumps({'available': True, 'node': str(node), 'artifact_tool': str(module)}, ensure_ascii=False))
            return 0
        if not args.project:
            parser.error('--project가 필요합니다 (런타임 확인은 --check-runtime).')
        report = render_project(args.project, args.view)
    except (RenderError, OSError, ValueError) as exc:
        print(f'내보내기 중단: {exc}', file=sys.stderr)
        return 2
    print(json.dumps({k: report[k] for k in ('status', 'slide_count', 'planned_seconds', 'font', 'warnings', 'outputs')}, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
