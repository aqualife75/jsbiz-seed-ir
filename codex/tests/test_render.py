"""Meaningful render contract, containment and editable-output regressions."""
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'plugins/seed-ir/skills/seed-ir/scripts/render_deck.py'
spec = importlib.util.spec_from_file_location('seed_ir_renderer', SCRIPT)
render = importlib.util.module_from_spec(spec)
spec.loader.exec_module(render)


def evidence():
    return {'items': [{'id': 'E1', 'claim': '교육용 가상 수치를 포함하는 검증 표',
                       'kind': 'internal', 'status': 'confirmed', 'source_id': 'training-fixture',
                       'locator': '1쪽', 'publisher': '가상 교육용 팀'}]}


def slide(sid='s1', layout='evidence'):
    return {'id': sid, 'section': 'problem', 'layout': layout,
            'governing_message': '반복 업무에 30분이 소요됩니다', 'governing_kind': 'fact',
            'governing_evidence_ids': ['E1'], 'evidence_ids': ['E1'],
            'body': [{'text': '담당자가 같은 데이터를 다시 입력합니다', 'kind': 'fact', 'evidence_ids': ['E1']}],
            'assets': [], 'seconds': 25, 'speaker_notes': '교육용 가상 사례입니다. 반복 입력을 줄이는 제품을 설명합니다.',
            'visual': {'type': 'none'}}


class TestRender(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        profile = json.loads((ROOT / 'plugins/seed-ir/skills/seed-ir/references/design-profile.json').read_text('utf-8'))
        try:
            cls.font = render.resolve_font(profile)
        except render.RenderError:
            cls.font = None

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='seed-ir-render-test-')
        self.addCleanup(self.temp.cleanup)
        self.project = Path(self.temp.name)
        self.profile = json.loads((ROOT / 'plugins/seed-ir/skills/seed-ir/references/design-profile.json').read_text('utf-8'))
        for folder in ('deck', 'evidence', 'design', 'assets'):
            (self.project / folder).mkdir()

    def save(self, deck, ledger=None):
        for relative, data in [('deck/deck.json', deck), ('evidence/evidence.json', ledger or evidence()),
                               ('design/design-profile.json', self.profile)]:
            (self.project / relative).write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')

    def test_all_layouts_editable_chart_notes_and_no_false_render_pass(self):
        if not self.font:
            self.skipTest('No supported Korean font installed')
        slides = [slide(f's{i}', layout) for i, layout in enumerate(sorted(render.LAYOUTS))]
        target = next(s for s in slides if s['layout'] == 'evidence')
        target['visual'] = {'type': 'bar', 'categories': ['기존', '실험'], 'values': [30, 10],
                            'unit': '분', 'source': '가상 교육용 측정표', 'evidence_ids': ['E1']}
        self.save({'team_name': '가상 교육용 팀', 'slides': slides})
        with patch.object(render, 'resolve_font', return_value=self.font):
            report = render.render_project(self.project)
        self.assertEqual(report['status'], 'draft_exported')
        self.assertEqual(report['pptx_render_review'], 'not_performed')
        self.assertFalse(report['storyboard_is_pptx_render'])
        self.assertEqual(report['editable_chart_count'], 1)
        deck = Presentation(str(self.project / 'output/seed-ir-draft.pptx'))
        self.assertEqual(len(deck.slides), len(render.LAYOUTS))
        self.assertAlmostEqual(deck.slide_width / 914400, 13 + 1 / 3, places=5)
        for actual in deck.slides:
            text = '\n'.join(shape.text for shape in actual.shapes if shape.has_text_frame)
            self.assertIn('30분', text)
            self.assertIn('E1', actual.notes_slide.notes_text_frame.text)
            self.assertIn('교육용 가상 사례', actual.notes_slide.notes_text_frame.text)
        story = (self.project / 'output/storyboard.html').read_text('utf-8')
        self.assertIn('PPTX를 렌더한 이미지가 아닙니다', story)
        with ZipFile(self.project / 'output/seed-ir-draft.pptx') as package:
            self.assertIn('ppt/charts/chart1.xml', package.namelist())
            self.assertTrue(any(name.startswith('ppt/embeddings/') for name in package.namelist()))
            chart_xml = ET.fromstring(package.read('ppt/charts/chart1.xml'))
            ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
            ids = chart_xml.findall('.//c:axId', ns) + chart_xml.findall('.//c:crossAx', ns)
            self.assertTrue(ids)
            self.assertTrue(all(0 <= int(x.attrib['val']) < 2**32 for x in ids))

    def test_fact_without_evidence_is_blocked(self):
        s = slide()
        s['body'][0]['evidence_ids'] = []
        with self.assertRaisesRegex(render.RenderError, '사실 주장'):
            render.validate_contract(self.project, {'team_name': 'T', 'slides': [s]}, evidence(), self.profile)

    def test_generic_rich_repeated_images_keep_crops_and_use_private_packaging_receipt(self):
        if not self.font:
            self.skipTest('No supported Korean font installed')
        try:
            render.artifact_runtime()
        except render.RenderError as error:
            self.skipTest(str(error))
        image_path = self.project / 'assets/repeated.png'
        Image.new('RGB', (120, 80), color=(190, 180, 170)).save(image_path)
        original_image = image_path.read_bytes()
        slides = []
        for sid, crop in [('custom-photo-a', None), ('custom-photo-b',
                {'left': .25, 'top': 0, 'right': .25, 'bottom': 0})]:
            current = slide(sid, 'evidence_board')
            current.update(section='product', blocks=[{
                'id': 'product-image', 'type': 'image', 'kind': 'assumption',
                'title': '가상 화면', 'evidence_ids': [],
                'path': 'assets/repeated.png', 'caption': '가상 설계 · 기능 검증용 이미지',
                'source': '자체 제작 테스트 fixture', 'rights': 'owned', 'alt': '원본 비트맵'},
                {'id': 'scope', 'type': 'text', 'kind': 'assumption', 'evidence_ids': [],
                 'title': '범위', 'text': '실제 제품이 아닌 기능 검증용 이미지입니다.'}])
            if crop:
                current['blocks'][0]['crop'] = crop
            current['visual_brief'] = {
                'key_message': current['governing_message'], 'primary_block_id': 'product-image',
                'reading_order': ['product-image', 'scope'], 'layout_variant': 'product_hero',
                'visual_reason': '같은 원본과 편집 가능한 크롭을 검증합니다.'}
            slides.append(current)
        self.save({'schema_version': '1.2', 'team_name': '비공유 기능검증 팀', 'slides': slides,
                   'views': {'master': {'slide_ids': [s['id'] for s in slides]},
                             'pitch': {'slide_ids': [s['id'] for s in slides]}}})
        result = render.render_project(self.project, 'master')
        self.assertEqual(result['slide_count'], 2)
        receipt = json.loads((self.project / 'work/render/master-media-dedup-report.json').read_text('utf-8'))
        self.assertEqual(receipt['media_count_after'], 1)
        self.assertTrue(receipt['preserved_media_identical'])
        self.assertFalse(Path(receipt['input']).is_relative_to(self.project))
        with ZipFile(self.project / 'output/master/seed-ir-draft.pptx') as package:
            media_parts = [name for name in package.namelist() if name.startswith('ppt/media/')]
            self.assertEqual(len(media_parts), 1)
            self.assertEqual(package.read(media_parts[0]), original_image)
            second = ET.fromstring(package.read('ppt/slides/slide2.xml'))
            crop = second.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect')
            self.assertEqual(crop.attrib, {'l': '25000', 't': '0', 'r': '25000', 'b': '0'})
        self.assertEqual(len(list((self.project / 'output/master/renders').glob('*.png'))), 2)

    def test_unverified_evidence_cannot_support_fact(self):
        e = evidence()
        e['items'][0]['status'] = 'unverified'
        with self.assertRaisesRegex(render.RenderError, '미검증'):
            render.validate_contract(self.project, {'team_name': 'T', 'slides': [slide()]}, e, self.profile)

    def test_body_over_capacity_not_silently_truncated(self):
        s = slide()
        s['body'] *= 4
        with self.assertRaisesRegex(render.RenderError, '최대 3개'):
            render.validate_contract(self.project, {'team_name': 'T', 'slides': [s]}, evidence(), self.profile)

    def test_unknown_image_rights_is_blocked(self):
        s = slide()
        s['assets'] = [{'path': 'assets/p.png', 'caption': '사진', 'source': 'url', 'rights': 'unknown', 'alt': '사진'}]
        with self.assertRaisesRegex(render.RenderError, 'unknown'):
            render.validate_contract(self.project, {'team_name': 'T', 'slides': [s]}, evidence(), self.profile)

    def test_image_outside_project_is_blocked(self):
        for value in ('../outside.png', '..\\outside.png', 'C:\\outside.png', '/etc/passwd', 'https://example.test/a.png'):
            with self.subTest(value=value), self.assertRaises(render.RenderError):
                render.contained_path(self.project, value)

    def test_licensed_image_is_embedded_with_caption_alt_and_source(self):
        if not self.font:
            self.skipTest('No supported Korean font installed')
        Image.new('RGB', (1200, 600), color=(190, 180, 170)).save(self.project / 'assets/p.png')
        s = slide(layout='cover')
        s['assets'] = [{'path': 'assets/p.png', 'caption': '사진 기능 검증용 이미지',
                        'source': '자체 제작 테스트 fixture', 'rights': 'owned', 'alt': '2대1 비율 테스트 이미지'}]
        self.save({'team_name': 'T', 'slides': [s]})
        with patch.object(render, 'resolve_font', return_value=self.font):
            render.render_project(self.project)
        deck = Presentation(str(self.project / 'output/seed-ir-draft.pptx'))
        pictures = [x for x in deck.slides[0].shapes if x.shape_type == 13]
        self.assertEqual(len(pictures), 1)
        self.assertAlmostEqual(pictures[0].width / pictures[0].height, 2.0, places=5)
        self.assertEqual(pictures[0]._element.nvPicPr.cNvPr.get('descr'), s['assets'][0]['alt'])
        self.assertIn('자체 제작 테스트 fixture', deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn('사진 기능 검증용 이미지', '\n'.join(x.text for x in deck.slides[0].shapes if x.has_text_frame))

    def test_chart_and_image_cannot_silently_overlap(self):
        Image.new('RGB', (1200, 600), color=(200, 200, 200)).save(self.project / 'assets/p.png')
        s = slide()
        s['assets'] = [{'path': 'assets/p.png', 'caption': '테스트 이미지', 'source': '테스트 자체 제작', 'rights': 'owned', 'alt': '회색 테스트 이미지'}]
        s['visual'] = {'type': 'bar', 'categories': ['A'], 'values': [10], 'unit': '명', 'source': '테스트', 'evidence_ids': ['E1']}
        with self.assertRaisesRegex(render.RenderError, '분리'):
            render.validate_contract(self.project, {'team_name': 'T', 'slides': [s]}, evidence(), self.profile)

    def test_unsupported_chart_layout_is_blocked(self):
        s = slide(layout='timeline')
        s['visual'] = {'type': 'bar', 'categories': ['A'], 'values': [10], 'unit': '명', 'source': '테스트', 'evidence_ids': ['E1']}
        with self.assertRaisesRegex(render.RenderError, 'evidence 또는 closing'):
            render.validate_contract(self.project, {'team_name': 'T', 'slides': [s]}, evidence(), self.profile)

    def test_chart_nonfinite_value_is_blocked(self):
        s = slide()
        s['visual'] = {'type': 'bar', 'categories': ['A'], 'values': [float('nan')], 'unit': '명', 'source': '테스트', 'evidence_ids': ['E1']}
        with self.assertRaisesRegex(render.RenderError, '유한 숫자'):
            render.validate_contract(self.project, {'team_name': 'T', 'slides': [s]}, evidence(), self.profile)

    def test_measured_overflow_preserves_existing_pptx(self):
        if not self.font:
            self.skipTest('No supported Korean font installed')
        s = slide(layout='process')
        s['body'] = [{'text': '가' * 80, 'kind': 'plan', 'evidence_ids': []} for _ in range(3)]
        for i in range(3):
            Image.new('RGB', (1200, 600), color=(200, 200, 200)).save(self.project / f'assets/p{i}.png')
            s['assets'].append({'path': f'assets/p{i}.png', 'caption': '테스트 이미지', 'source': '직접 생성', 'rights': 'owned', 'alt': '테스트'})
        self.save({'team_name': 'T', 'slides': [s]})
        (self.project / 'output').mkdir()
        candidate = self.project / 'output/seed-ir-draft.pptx'
        candidate.write_bytes(b'prior-output')
        with patch.object(render, 'resolve_font', return_value=self.font):
            with self.assertRaisesRegex(render.RenderError, '영역'):
                render.render_project(self.project)
        self.assertEqual(candidate.read_bytes(), b'prior-output')

    def test_html_escapes_source_and_body(self):
        s = slide()
        s['body'][0]['text'] = '<script>alert(1)</script>'
        target = self.project / 'story.html'
        render.write_storyboard(self.project, {'team_name': '<test>', 'slides': [s]}, {'E1': evidence()['items'][0]}, target)
        actual = target.read_text('utf-8')
        self.assertNotIn('<script>alert', actual)
        self.assertIn('&lt;script&gt;', actual)


if __name__ == '__main__':
    unittest.main()
