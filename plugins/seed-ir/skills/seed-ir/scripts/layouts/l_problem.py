# -*- coding: utf-8 -*-
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from . import register
from .base import rgb, _set_font
import design_system as ds

_TONE = {"accent": ds.C["accent"], "amber": ds.C["amber"], "teal": ds.C["teal"], "purple": ds.C["purple"]}
_PP_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}

def _chrome(cv, s, slide, lead_w=ds.CW):
    cv.kicker(s.get("kicker", "")); cv.title(s.get("title", ""), w=lead_w); cv.lead(s.get("lead", ""), w=lead_w)
    cv.source_line(slide.get("source_line", "")); cv.page_no(slide["no"], cv.ctx.total); cv.notes(slide.get("notes"))

@register("problem_cascade")
def problem_cascade(cv, s, slide):
    has_panel = bool(s.get("panel_image")) or bool(s.get("quote"))
    left_w = (ds.PANEL_X - ds.MX - 0.6) if has_panel else ds.CW
    if s.get("panel_image"): cv.panel_image(s["panel_image"], strength=0.65)
    _chrome(cv, s, slide, lead_w=left_w)
    bars = s.get("bars") or []
    cv.card(ds.MX, ds.BODY_Y, left_w, 3.0)
    cv.label(ds.MX + 0.4, ds.BODY_Y + 0.35, left_w - 0.8, "CASCADE · 단계별 이탈")
    if bars:
        mx = max(b["value"] for b in bars) or 1; bw = ds.col_w(len(bars), left_w - 0.8, 0.2); bx = ds.MX + 0.4; base_y = ds.BODY_Y + 2.35
        for i, b in enumerate(bars):
            x = bx + i * (bw + 0.2); h = 0.9 * b["value"] / mx; tone = _TONE.get(b.get("tone", ["teal", "amber", "accent", "accent"][i % 4]))
            cv.text(x, base_y - h - 0.5, bw, 0.45, f"{b['value']}%", 24, tone if b.get("tone") == "accent" else cv.ink, bold=True, max_lines=1, slot="bars.value")
            cv.rect(x, base_y - h, bw, h, fill=tone)
            cv.text(x, base_y + 0.08, bw, 0.3, b["label"], ds.SZ["body_sm"], cv.ink, bold=True, max_lines=1, slot="bars.label")
    kpis = s.get("kpis") or []
    if kpis:
        kw = ds.col_w(len(kpis), left_w, 0.2); ky = ds.BODY_Y + 3.35; kh = ds.BODY_B - ky
        for i, k in enumerate(kpis):
            x = ds.MX + i * (kw + 0.2); tone = list(_TONE.values())[i % 3]
            cv.card(x, ky, kw, kh); cv.rect(x, ky, 0.05, kh, fill=tone)
            cv.stat(x + 0.35, ky + 0.35, kw - 0.6, k.get("big", ""), label=k.get("label", ""), source=k.get("source"), big_size=ds.SZ["big_md"])
    q = s.get("quote")
    if q:
        qx, qy, qw, qh = ds.PANEL_X + 0.4, 4.5, ds.PANEL_W - 0.8, 3.3
        cv.rect(qx, qy, qw, qh, fill=ds.BG["dark"], alpha=70, line=ds.C["amber"], line_alpha=60, radius=True)
        cv.label(qx + 0.4, qy + 0.35, qw - 0.8, "현장의 언어", color=ds.C["amber"])
        cv.text(qx + 0.4, qy + 0.8, qw - 0.8, 1.9, f"“{q.get('text','')}”", ds.SZ["card_title"], ds.C["white"], bold=True, line_spacing=1.3, max_lines=4, slot="quote.text")
        cv.text(qx + 0.4, qy + 2.75, qw - 0.8, 0.5, q.get("who", ""), ds.SZ["small"], ds.C["sub_dark"], max_lines=2, slot="quote.who")

@register("problem_grid")
def problem_grid(cv, s, slide):
    _chrome(cv, s, slide)
    stats = s.get("stats") or []
    y = ds.BODY_Y
    if stats:
        cv.label(ds.MX, y, 8, "EVIDENCE · 이탈률")
        sw = ds.col_w(len(stats));
        for i, st in enumerate(stats):
            cv.stat(ds.MX + i * (sw + ds.GAP), y + 0.4, sw - 0.3, st.get("big", ""), label=st.get("label", ""), source=st.get("source"), big_size=ds.SZ["big_md"], color=cv.accent if i == 0 else cv.ink)
        y += 2.0
    pts = s.get("points") or []
    if pts:
        pw = ds.col_w(len(pts)); ph = ds.BODY_B - y - 0.5
        for i, p in enumerate(pts):
            x = ds.MX + i * (pw + ds.GAP); cv.card(x, y, pw, ph)
            cv.text(x + 0.35, y + 0.3, pw - 0.7, 0.5, p.get("num", f"0{i+1}"), 30, cv.accent, bold=True, max_lines=1, slot="points.num")
            cv.text(x + 0.35, y + 0.95, pw - 0.7, 0.5, p.get("title", ""), ds.SZ["card_title"], cv.ink, bold=True, max_lines=1, slot="points.title")
            cv.text(x + 0.35, y + 1.45, pw - 0.7, 0.4, p.get("sub", ""), ds.SZ["body"], cv.accent, bold=True, max_lines=1, slot="points.sub")
            cv.text(x + 0.35, y + 1.95, pw - 0.7, ph - 2.2, p.get("desc", ""), ds.SZ["body"], cv.sub, line_spacing=1.45, max_lines=4, slot="points.desc")
    if s.get("footnote"): cv.text(ds.MX, ds.BODY_B - 0.4, ds.CW, 0.35, s["footnote"], ds.SZ["small"], cv.muted, max_lines=1, slot="footnote")

@register("alt_table")
def alt_table(cv, s, slide):
    _chrome(cv, s, slide)
    t = s.get("table") or {}; cols = t.get("columns") or []; rows = t.get("rows") or []; us = t.get("us_col", len(cols) - 1)
    n_r, n_c = len(rows) + 1, len(cols) + 1
    th = min(0.85, (ds.BODY_B - ds.BODY_Y - 1.2) / n_r)
    gf = cv.s.shapes.add_table(n_r, n_c, Inches(ds.MX), Inches(ds.BODY_Y), Inches(ds.CW), Inches(th * n_r)); tbl = gf.table
    tbl.columns[0].width = Inches(3.6)
    for j in range(1, n_c): tbl.columns[j].width = Inches((ds.CW - 3.6) / len(cols))
    def cell(r, c, text, size, color, bold=False, fill=None, align="left"):
        ce = tbl.cell(r, c); ce.text = ""; ce.margin_left = ce.margin_right = Inches(0.2); ce.margin_top = ce.margin_bottom = Inches(0.08)
        ce.fill.solid(); ce.fill.fore_color.rgb = rgb(fill or ("1A2130" if cv.bg == "dark" else "FFFFFF"))
        p = ce.text_frame.paragraphs[0]; p.alignment = _PP_ALIGN[align]; run = p.add_run(); run.text = text; _set_font(run, ds.FONT, size, bold, color)
    cell(0, 0, "비교 항목", ds.SZ["body_sm"], ds.C["amber"], True, fill=ds.BG["dark"])
    for j, c in enumerate(cols, 1):
        cell(0, j, c, ds.SZ["body_sm"], ds.C["white"] if j - 1 != us else ds.C["accent_dark"], True, fill=ds.BG["dark"], align="center")
    for i, r in enumerate(rows, 1):
        weak = r.get("weakness", False)
        cell(i, 0, r.get("label", ""), ds.SZ["body"], cv.ink, True, fill=ds.C["amber_tint"] if weak else None)
        for j, v in enumerate(r.get("cells", []), 1):
            is_us = (j - 1 == us)
            color = (ds.C["amber"] if weak else cv.accent) if is_us else cv.sub
            cell(i, j, v, ds.SZ["body"], color, is_us, fill=(ds.C["amber_tint"] if weak else (ds.C["accent_tint"] if is_us and cv.bg == "cream" else None)), align="center")
    if any(r.get("weakness") for r in rows):
        cv.text(ds.MX, ds.BODY_Y + th * n_r + 0.15, ds.CW, 0.3, "▲ 강조 행 = 우리가 뒤지는 항목(솔직하게 표시)", ds.SZ["small"], ds.C["amber"], bold=True, max_lines=1)
    if s.get("premise_note"):
        cv.card(ds.MX, ds.BODY_B - 1.0, ds.CW, 0.85, style="tint" if cv.bg == "cream" else "plain")
        cv.text(ds.MX + 0.4, ds.BODY_B - 0.9, ds.CW - 0.8, 0.65, s["premise_note"], ds.SZ["body"], cv.ink, anchor="middle", max_lines=2, slot="premise_note")

@register("quadrant")
def quadrant(cv, s, slide):
    _chrome(cv, s, slide)
    ax = s.get("axes") or {}
    qx, qy, qw, qh = ds.MX, ds.BODY_Y, 10.4, ds.BODY_B - ds.BODY_Y - 0.4
    line = ds.C["muted_dark"] if cv.bg == "dark" else ds.C["card_border"]
    cv.rect(qx, qy, 0.02, qh, fill=line); cv.rect(qx, qy + qh, qw, 0.02, fill=line)
    cv.rect(qx + qw / 2, qy, 0.015, qh, fill=line, alpha=50); cv.rect(qx, qy + qh / 2, qw, 0.015, fill=line, alpha=50)
    cv.text(qx + 0.2, qy + 0.1, 4, 0.3, ax.get("x_left", ""), ds.SZ["small"], cv.muted, max_lines=1); cv.text(qx + qw - 4.2, qy + 0.1, 4, 0.3, ax.get("x_right", ""), ds.SZ["small"], cv.muted, align="right", max_lines=1)
    cv.text(qx - 1.0, qy + qh / 2 - 0.5, 1.0, 0.9, f"{ax.get('y_top','')}\n{ax.get('y_bottom','')}", ds.SZ["small"], cv.muted, align="right")
    for b in s.get("bubbles") or []:
        d = b.get("size", 1.6); cx = qx + b["x"] * qw - d / 2; cy = qy + (1 - b["y"]) * qh - d / 2
        us = b.get("is_us"); tone = _TONE.get(b.get("tone", ""), ds.C["muted_dark"] if cv.bg == "dark" else ds.C["sub"])
        shp = cv.s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d)); shp.shadow.inherit = False
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(cv.accent if us else tone)
        if not us:
            from .base import _alpha; from pptx.oxml.ns import qn
            _alpha(shp.fill._xPr.find(qn("a:solidFill")), 22); shp.line.color.rgb = rgb(tone); shp.line.width = Pt(1)
        else: shp.line.fill.background()
        txt_color = ds.C["white"] if us else (ds.C["white"] if cv.bg == "dark" else cv.ink)
        cv.text(cx, cy + d * 0.28, d, 0.45, b.get("name", ""), 17 if us else 14, txt_color, bold=True, align="center", max_lines=1, slot="bubbles.name")
        cv.text(cx, cy + d * 0.28 + 0.45, d, 0.5, b.get("sub", ""), ds.SZ["small"], txt_color if us else cv.muted, align="center", max_lines=2, slot="bubbles.sub")
    gaps = s.get("gaps") or []
    gx, gw = ds.MX + qw + 0.7, ds.CW - qw - 0.7
    cv.label(gx, ds.BODY_Y, gw, "대체재가 남긴 공백", color=ds.C["amber"])
    gh = (ds.BODY_B - ds.BODY_Y - 0.6 - 0.2 * (len(gaps) - 1)) / max(1, len(gaps))
    for i, g in enumerate(gaps):
        y = ds.BODY_Y + 0.45 + i * (gh + 0.2); cv.card(gx, y, gw, gh); cv.rect(gx, y, 0.05, gh, fill=list(_TONE.values())[i % 3])
        cv.text(gx + 0.35, y + 0.25, gw - 0.6, 0.45, g.get("title", ""), ds.SZ["card_title"], cv.ink, bold=True, max_lines=1, slot="gaps.title")
        cv.text(gx + 0.35, y + 0.75, gw - 0.6, gh - 0.9, g.get("desc", ""), ds.SZ["body"], cv.sub, line_spacing=1.4, max_lines=3, slot="gaps.desc")
