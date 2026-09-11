# -*- coding: utf-8 -*-
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from . import register
from .base import rgb, _set_font
from .l_problem import _chrome, _TONE
import design_system as ds

def _big_with_unit(cv, x, y, w, big, unit, size, color, align="left", slot=""):
    """큰 숫자 + 작은 단위 run (market_tam 계단 박스용)."""
    tb = cv.text(x, y, w, size / 72 * 1.3, big, size, color, bold=True, align=align, max_lines=1, slot=slot)
    if unit:
        r = tb.text_frame.paragraphs[0].add_run(); r.text = unit; _set_font(r, ds.FONT, ds.SZ["unit"], True, color)
    return tb

@register("traction_plan")
def traction_plan(cv, s, slide):
    _chrome(cv, s, slide)
    weeks = s.get("weeks") or []; y = ds.BODY_Y
    cv.label(ds.MX, y, ds.CW, "행동 기반 검증 설계", color=cv.accent)
    if weeks:
        w = ds.col_w(len(weeks)); h = 2.35
        for i, wk in enumerate(weeks):
            x = ds.MX + i * (w + ds.GAP); dark = wk.get("tone") == "dark"; tone = list(_TONE.values())[i % 4]
            cv.card(x, y + 0.4, w, h, style="dark" if dark else "plain", top_bar=tone)
            ink, sub = (ds.C["white"], ds.C["sub_dark"]) if dark else (cv.ink, cv.sub)
            cv.label(x + 0.35, y + 0.75, w - 0.7, wk.get("label", ""), color=ds.C["teal"] if dark else tone)
            cv.text(x + 0.35, y + 1.1, w - 0.7, 0.55, wk.get("big", ""), 24, ink, bold=True, max_lines=1, slot="weeks.big")
            cv.text(x + 0.35, y + 1.7, w - 0.7, 1.0, wk.get("desc", ""), ds.SZ["body_sm"], sub, line_spacing=1.4, max_lines=3, slot="weeks.desc")
        y += h + 0.75
    rows = (s.get("table") or {}).get("rows") or []
    lw = 11.0
    if rows:
        cv.label(ds.MX, y, lw, "진행 판단 기준 (Go / Pivot Signal)", color=cv.accent)
        n_r = len(rows) + 1; th = min(0.6, (ds.BODY_B - y - 0.5) / n_r)
        tbl = cv.s.shapes.add_table(n_r, 3, Inches(ds.MX), Inches(y + 0.4), Inches(lw), Inches(th * n_r)).table
        tbl.columns[0].width = Inches(3.6); tbl.columns[1].width = Inches(2.0); tbl.columns[2].width = Inches(lw - 5.6)
        from .base import _set_font
        def cell(r, c, t, color, bold=False, fill=None):
            ce = tbl.cell(r, c); ce.text = ""; ce.margin_top = ce.margin_bottom = Inches(0.05); ce.fill.solid(); ce.fill.fore_color.rgb = rgb(fill or ("1A2130" if cv.bg == "dark" else "FFFFFF"))
            run = ce.text_frame.paragraphs[0].add_run(); run.text = t; _set_font(run, ds.FONT, ds.SZ["body_sm"], bold, color)
        for j, hd in enumerate(["지표", "목표", "측정 방법 — 말이 아닌 행동"]): cell(0, j, hd, ds.C["amber"] if j == 0 else ds.C["white"], True, fill=ds.BG["dark"])
        for i, r in enumerate(rows, 1):
            cell(i, 0, r.get("metric", ""), cv.ink, True); cell(i, 1, r.get("target", ""), cv.accent, True); cell(i, 2, r.get("method", ""), cv.sub)
    side = s.get("side") or {}; sx = ds.MX + lw + 0.35; sw = ds.CW - lw - 0.35
    sy = y + 0.4; sh = 1.5
    if side:
        cv.card(sx, sy, sw, sh); cv.label(sx + 0.35, sy + 0.25, sw - 0.7, side.get("title", ""), color=cv.accent)
        cv.text(sx + 0.35, sy + 0.65, sw - 0.7, sh - 0.85, side.get("body", ""), ds.SZ["body_sm"], cv.sub, line_spacing=1.4, max_lines=3, slot="side.body")
    if s.get("insight"):
        iy = sy + sh + 0.2; ih = ds.BODY_B - iy
        if ih > 0.6:
            cv.card(sx, iy, sw, ih, style="accent"); cv.label(sx + 0.35, iy + 0.22, sw - 0.7, "이미 확인된 것", color="FFFFFF")
            cv.text(sx + 0.35, iy + 0.5, sw - 0.7, ih - 0.6, s["insight"], ds.SZ["body_sm"], ds.C["white"], bold=True, line_spacing=1.3, max_lines=2, slot="insight")

@register("kpi_chart")
def kpi_chart(cv, s, slide):
    _chrome(cv, s, slide)
    ch = s.get("chart") or {}; cats = ch.get("categories") or []; series = [x for x in (ch.get("series") or []) if x.get("values")]
    cw = 11.0; y = ds.BODY_Y; h = ds.BODY_B - y - 0.7
    if cats and series:
        cd = CategoryChartData(); cd.categories = cats
        for se in series: cd.add_series(se["name"], se["values"])
        ctype = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS}.get(ch.get("type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        cv.card(ds.MX, y, cw, h)
        gf = cv.s.shapes.add_chart(ctype, Inches(ds.MX + 0.3), Inches(y + 0.3), Inches(cw - 0.6), Inches(h - 0.6), cd); chart = gf.chart
        chart.has_legend = len(series) > 1; chart.has_title = False
        chart.font.name = ds.FONT; chart.font.size = Pt(12); chart.font.color.rgb = rgb(cv.sub)
        plot = chart.plots[0]; plot.has_data_labels = True; plot.data_labels.font.size = Pt(13); plot.data_labels.font.bold = True; plot.data_labels.font.color.rgb = rgb(cv.ink)
        plot.data_labels.number_format = '#,##0'; plot.data_labels.number_format_is_linked = False
        if ctype != XL_CHART_TYPE.LINE_MARKERS: plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; plot.gap_width = 80
        for k, se in enumerate(plot.series):
            se.format.fill.solid(); se.format.fill.fore_color.rgb = rgb([cv.accent, ds.C["amber"], ds.C["teal"]][k % 3])
        chart.value_axis.has_major_gridlines = False; chart.value_axis.visible = False
        chart.category_axis.format.line.color.rgb = rgb(ds.C["card_border"]); chart.category_axis.tick_labels.font.size = Pt(12)
    else:
        cv.ctx.warnings.append({"slide": cv.ctx.current_no, "slot": "chart", "text": "값 없는 차트 — 카드만 표시", "lines": 0, "max": 0})
    kpis = s.get("kpis") or []; kx = ds.MX + cw + 0.35; kw = ds.CW - cw - 0.35; kh = (h - 0.2 * 2) / 3
    for i, k in enumerate(kpis[:3]):
        ky = y + i * (kh + 0.2); cv.card(kx, ky, kw, kh); cv.rect(kx, ky, 0.05, kh, fill=list(_TONE.values())[i % 3])
        cv.stat(kx + 0.35, ky + 0.3, kw - 0.6, k.get("big", ""), label=k.get("label", ""), source=k.get("source"), big_size=ds.SZ["big_md"], color=cv.accent if i == 0 else cv.ink)
    if s.get("note"): cv.text(ds.MX, ds.BODY_B - 0.45, ds.CW, 0.35, s["note"], ds.SZ["small"], cv.muted, max_lines=1, slot="note")

@register("bm_pricing")
def bm_pricing(cv, s, slide):
    _chrome(cv, s, slide)
    tiers = s.get("tiers") or []; y = ds.BODY_Y; h = 3.7
    if tiers:
        w = ds.col_w(len(tiers))
        for i, t in enumerate(tiers):
            x = ds.MX + i * (w + ds.GAP); hi = t.get("highlight")
            cv.card(x, y, w, h, style="accent" if hi else ("plain" if cv.bg == "cream" else "plain"))
            if not hi and i == len(tiers) - 1 and cv.bg == "dark": cv.rect(x, y, w, h, fill=None, line=ds.C["teal"], line_alpha=60, radius=True)
            ink, sub = (ds.C["white"], ds.C["white"]) if hi else (cv.ink, cv.sub)
            cv.label(x + 0.4, y + 0.35, w - 0.8, t.get("label", ""), color=ds.C["white"] if hi else (ds.C["teal"] if i == len(tiers) - 1 else cv.accent))
            cv.text(x + 0.4, y + 0.7, w - 0.8, 0.75, t.get("price", ""), 34, ink, bold=True, max_lines=1, slot="tiers.price")
            cv.text(x + 0.4, y + 1.45, w - 0.8, 0.35, t.get("sub", ""), ds.SZ["body_sm"], sub, max_lines=1, slot="tiers.sub")
            cv.body_bullets(x + 0.4, y + 1.9, w - 0.8, 1.1, t.get("bullets") or [], size=ds.SZ["body_sm"], color=ink if hi else cv.ink)
            cv.text(x + 0.4, y + 3.0, w - 0.8, 0.6, t.get("note", ""), ds.SZ["small"], sub, bold=hi, line_spacing=1.3, max_lines=2, slot="tiers.note")
        y += h + 0.35
    ue = s.get("unit_econ") or []
    lw = 11.0
    if ue:
        cv.label(ds.MX, y, lw, "유닛 이코노믹스 가정", color=ds.C["amber"])
        uh = ds.BODY_B - y - 0.45; cv.card(ds.MX, y + 0.4, lw, uh); uw = (lw - 0.8) / len(ue)
        for i, u in enumerate(ue):
            ux = ds.MX + 0.4 + i * uw
            cv.label(ux, y + 0.7, uw - 0.2, u.get("label", ""), color=[cv.ink, cv.ink, ds.C["teal"], ds.C["amber"]][i % 4] if i >= 2 else cv.sub)
            cv.text(ux, y + 1.05, uw - 0.2, 0.55, u.get("big", ""), 24, [cv.ink, cv.ink, ds.C["teal"], ds.C["amber"]][i % 4], bold=True, max_lines=1, slot="unit_econ.big")
            cv.text(ux, y + 1.6, uw - 0.2, 0.35, u.get("sub", ""), ds.SZ["small"], cv.muted, max_lines=1, slot="unit_econ.sub")
    if s.get("evidence_note"):
        ex = ds.MX + lw + 0.35; ew = ds.CW - lw - 0.35
        cv.rect(ex, y + 0.4, ew, ds.BODY_B - y - 0.45, fill=None, line=ds.C["amber"], line_alpha=70, radius=True)
        cv.label(ex + 0.35, y + 0.7, ew - 0.7, "지불 의향 근거", color=ds.C["amber"])
        cv.text(ex + 0.35, y + 1.05, ew - 0.7, ds.BODY_B - y - 1.2, s["evidence_note"], ds.SZ["body_sm"], cv.sub, line_spacing=1.4, max_lines=4, slot="evidence_note")

@register("market_tam")
def market_tam(cv, s, slide):
    _chrome(cv, s, slide, lead_w=11.5)
    lw = 11.0; y = ds.BODY_Y
    boxes = [("tam", 0.0, "tint"), ("sam", 0.75, "plain"), ("som", 1.5, "dark")]
    bh = 1.45
    for i, (key, indent, style) in enumerate(boxes):
        b = s.get(key) or {}; x = ds.MX + indent; w = lw - indent * 2; by = y + i * (bh + 0.2)
        cv.card(x, by, w, bh, style=style if cv.bg == "cream" else ("dark" if key == "som" else "plain"))
        dark = style == "dark" or cv.bg == "dark"; ink, sub, mut = (ds.C["white"], ds.C["sub_dark"], ds.C["muted_dark"]) if dark else (cv.ink, cv.sub, cv.muted)
        cv.label(x + 0.4, by + 0.25, w - 4.4, b.get("label", ""), color=cv.accent if not dark else ds.C["amber"])
        cv.text(x + 0.4, by + 0.6, w - 4.4, 0.75, b.get("desc", ""), ds.SZ["body"], ink, bold=True, line_spacing=1.3, max_lines=2, slot=f"{key}.desc")
        _big_with_unit(cv, x + w - 4.0, by + 0.2, 3.6, b.get("big", ""), b.get("unit", ""), ds.SZ["big_md"], ink, align="right", slot=f"{key}.big")
        cv.text(x + w - 4.0, by + 0.95, 3.6, 0.35, b.get("calc", ""), ds.SZ["small"], mut, align="right", max_lines=1, slot=f"{key}.calc")
    if s.get("sanity_note"):
        sy = y + 3 * (bh + 0.2) + 0.1; sh = ds.BODY_B - sy
        cv.rect(ds.MX, sy, lw, sh, fill=("E8F3EE" if cv.bg == "cream" else "FFFFFF"), alpha=(None if cv.bg == "cream" else 6), line=ds.C["teal_deep"], line_alpha=50, radius=True)
        cv.text(ds.MX + 0.4, sy + 0.15, lw - 0.8, sh - 0.3, "**기준 검증** — " + s["sanity_note"], ds.SZ["body"], cv.ink, bold=False, anchor="middle", line_spacing=1.4, max_lines=3, slot="sanity_note")
    st = s.get("side_stats") or []; sx = ds.MX + lw + 0.6; sw = ds.CW - lw - 0.6
    cv.label(sx, y, sw, "시장은 커지는 중", color=cv.accent)
    sh_ = (ds.BODY_B - y - 0.45 - 0.2 * 2) / 3
    for i, it in enumerate(st[:3]):
        iy = y + 0.45 + i * (sh_ + 0.2); cv.card(sx, iy, sw, sh_)
        cv.label(sx + 0.35, iy + 0.25, sw - 0.6, it.get("label", ""), color=[cv.accent, ds.C["amber"], cv.accent][i % 3])
        cv.text(sx + 0.35, iy + 0.6, sw - 0.6, 0.6, it.get("big", ""), 26, cv.ink, bold=True, max_lines=1, slot="side_stats.big")
        cv.text(sx + 0.35, iy + 1.2, sw - 0.6, 0.35, it.get("sub", ""), ds.SZ["body_sm"], cv.sub, max_lines=1, slot="side_stats.sub")
        cv.text(sx + 0.35, iy + sh_ - 0.4, sw - 0.6, 0.3, it.get("source", ""), ds.SZ["source"], cv.muted, max_lines=1, slot="side_stats.source")

@register("gtm_funnel")
def gtm_funnel(cv, s, slide):
    _chrome(cv, s, slide)
    bh = s.get("beachhead") or {}; y = ds.BODY_Y; lw = 8.6; h = 3.3
    cv.card(ds.MX, y, lw, h, style="dark"); cv.label(ds.MX + 0.4, y + 0.35, lw - 0.8, "BEACHHEAD · 교두보 고객", color=ds.C["teal"])
    cv.text(ds.MX + 0.4, y + 0.75, lw - 0.8, 1.2, bh.get("title", ""), 26, ds.C["white"], bold=True, line_spacing=1.25, max_lines=2, slot="beachhead.title")
    cv.rect(ds.MX + 0.4, y + 2.0, lw - 0.8, 0.01, fill="FFFFFF", alpha=15)
    half = (lw - 1.0) / 2
    cv.text(ds.MX + 0.4, y + 2.1, half, 0.3, "왜 이 고객군인가", ds.SZ["body_sm"], ds.C["teal"], bold=True, max_lines=1)
    cv.text(ds.MX + 0.4, y + 2.4, half, 0.85, bh.get("why", ""), ds.SZ["small"], ds.C["sub_dark"], line_spacing=1.35, max_lines=3, slot="beachhead.why")
    cv.text(ds.MX + 0.6 + half, y + 2.1, half, 0.3, "초기 제외 대상", ds.SZ["body_sm"], ds.C["amber"], bold=True, max_lines=1)
    cv.text(ds.MX + 0.6 + half, y + 2.4, half, 0.85, bh.get("exclude", ""), ds.SZ["small"], ds.C["sub_dark"], line_spacing=1.35, max_lines=3, slot="beachhead.exclude")
    chs = s.get("channels") or []; cx = ds.MX + lw + 0.35; cw = ds.CW - lw - 0.35
    cv.card(cx, y, cw, h); cv.label(cx + 0.4, y + 0.35, cw - 0.8, "채널 우선순위 — 신뢰가 CAC를 결정합니다")
    for i, c in enumerate(chs[:3]):
        cy = y + 0.8 + i * 0.8
        cv.text(cx + 0.4, cy, cw - 0.8, 0.32, c.get("title", ""), ds.SZ["body"], cv.ink, bold=True, max_lines=1, slot="channels.title")
        cv.text(cx + 0.4, cy + 0.32, cw - 0.8, 0.45, c.get("desc", ""), ds.SZ["small"], cv.sub, max_lines=2, slot="channels.desc")
    fun = s.get("funnel") or []; fy = y + h + 0.4
    cv.label(ds.MX, fy, ds.CW, "확보 퍼널", color=cv.accent)
    if fun:
        fw = ds.col_w(len(fun)); fh = ds.BODY_B - fy - 0.85
        for i, f in enumerate(fun):
            fx = ds.MX + i * (fw + ds.GAP); dark = f.get("tone") == "dark"
            cv.card(fx, fy + 0.4, fw, fh, style="dark" if dark else "plain", top_bar=list(_TONE.values())[i % 4])
            ink, sub = (ds.C["white"], ds.C["sub_dark"]) if dark else (cv.ink, cv.sub)
            cv.label(fx + 0.3, fy + 0.7, fw - 0.6, f.get("label", ""), color=ds.C["accent_dark"] if dark else cv.accent)
            cv.stat(fx + 0.3, fy + 1.05, fw - 0.6, f.get("big", ""), f.get("unit", ""), big_size=ds.SZ["big_sm"], color=ink)
            cv.text(fx + 0.3, fy + 1.75, fw - 0.6, fh - 1.5, f.get("desc", ""), ds.SZ["small"], sub, max_lines=2, slot="funnel.desc")
    if s.get("footnote"): cv.text(ds.MX, ds.BODY_B - 0.35, ds.CW, 0.3, s["footnote"], ds.SZ["small"], cv.muted, max_lines=1, slot="footnote")

@register("growth_phases")
def growth_phases(cv, s, slide):
    _chrome(cv, s, slide)
    ph = s.get("phases") or []; y = ds.BODY_Y; h = 3.5
    if ph:
        w = ds.col_w(len(ph))
        for i, p in enumerate(ph):
            x = ds.MX + i * (w + ds.GAP); cv.card(x, y, w, h, top_bar=list(_TONE.values())[i % 3])
            cv.label(x + 0.4, y + 0.35, w - 0.8, p.get("label", ""), color=list(_TONE.values())[i % 3])
            cv.text(x + 0.4, y + 0.75, w - 0.8, 0.55, p.get("title", ""), 24, cv.ink, bold=True, max_lines=1, slot="phases.title")
            cv.text(x + 0.4, y + 1.35, w - 0.8, 1.2, p.get("desc", ""), ds.SZ["body"], cv.sub, line_spacing=1.4, max_lines=3, slot="phases.desc")
            cv.rect(x + 0.4, y + 2.6, w - 0.8, 0.01, fill=ds.C["card_border"])
            cv.text(x + 0.4, y + 2.7, w - 0.8, 0.3, "핵심 지표", ds.SZ["small"], ds.C["teal_deep"], bold=True, max_lines=1)
            cv.text(x + 0.4, y + 3.0, w - 0.8, 0.4, p.get("kpi", ""), ds.SZ["body_sm"], cv.ink, bold=True, max_lines=1, slot="phases.kpi")
        y += h + 0.4
    fw_items = s.get("flywheel") or []; lw = 11.0
    if fw_items:
        cv.card(ds.MX, y, lw, ds.BODY_B - y, style="dark" if cv.bg == "cream" else "plain"); cv.label(ds.MX + 0.4, y + 0.3, lw, "데이터 플라이휠", color=ds.C["teal"])
        n = len(fw_items); iw = (lw - 0.8) / n
        for i, it in enumerate(fw_items):
            cv.text(ds.MX + 0.4 + i * iw, y + 0.75, iw - 0.5, ds.BODY_B - y - 0.9, it, ds.SZ["body"], ds.C["white"], bold=True, anchor="middle", line_spacing=1.3, max_lines=3, slot="flywheel")
            if i < n - 1: cv.text(ds.MX + 0.4 + (i + 1) * iw - 0.5, y + 0.75, 0.5, ds.BODY_B - y - 0.9, "→", 22, ds.C["amber"], bold=True, align="center", anchor="middle")
    if s.get("expansion"):
        ex = ds.MX + lw + 0.35; ew = ds.CW - lw - 0.35
        cv.card(ex, y, ew, ds.BODY_B - y, style="tint" if cv.bg == "cream" else "plain"); cv.label(ex + 0.35, y + 0.3, ew - 0.7, "확장 옵션 · 같은 엔진, 다른 고객")
        cv.text(ex + 0.35, y + 0.7, ew - 0.7, ds.BODY_B - y - 0.85, s["expansion"], ds.SZ["body_sm"], cv.sub, line_spacing=1.4, max_lines=4, slot="expansion")
