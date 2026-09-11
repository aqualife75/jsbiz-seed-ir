# -*- coding: utf-8 -*-
"""Canvas 프리미티브 — 모든 레이아웃이 이것만으로 그린다."""
from __future__ import annotations
import math, os, re, tempfile
from dataclasses import dataclass, field
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
import design_system as ds
import prep_image

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_MARK = re.compile(r"(\*\*.+?\*\*|__.+?__)")

@dataclass
class DeckCtx:
    team: str
    total: int
    assets_dir: Path
    accent_override: str | None = None
    tmp: Path = field(default_factory=lambda: Path(tempfile.mkdtemp(prefix="seed_ir_")))
    warnings: list = field(default_factory=list)
    current_no: int = 0

def rgb(h: str) -> RGBColor: return RGBColor.from_string(h.lstrip("#").upper())

def est_lines(text: str, size_pt: float, width_in: float) -> int:
    """문단별 예상 줄 수 합. 한글 0.92em·ASCII 0.55em 평균 폭."""
    total = 0
    for para in (text or "").split("\n"):
        if not para: total += 1; continue
        w = sum(0.55 if ord(ch) < 0x2E80 else 0.92 for ch in para) * size_pt / 72.0
        total += max(1, math.ceil(w / max(width_in - 0.1, 0.1)))
    return total

def _set_font(run, name, size, bold, color, spacing=0.0):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)
    if spacing:
        rPr.set("spc", str(int(spacing * size * 100)))

def _alpha(color_elm_parent, pct: int):
    clr = color_elm_parent.find(qn("a:srgbClr"))
    if clr is not None:
        a = clr.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))}); clr.append(a)

class Canvas:
    def __init__(self, slide, bg: str, ctx: DeckCtx):
        self.s, self.bg, self.ctx = slide, bg, ctx
        self.accent = ctx.accent_override if (ctx.accent_override and bg == "cream") else ds.accent(bg)
        self.ink, self.sub, self.muted = ds.ink(bg), ds.sub(bg), ds.muted(bg)
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(ds.BG[bg])

    # ── 도형 ──
    def rect(self, x, y, w, h, fill=None, alpha=None, line=None, line_alpha=None, radius=False, gradient_to=None):
        shp = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius: shp.adjustments[0] = ds.RADIUS_ADJ if radius is True else radius
        shp.shadow.inherit = False
        if fill is None: shp.fill.background()
        elif gradient_to:
            shp.fill.gradient(); shp.fill.gradient_angle = 0
            shp.fill.gradient_stops[0].color.rgb = rgb(fill); shp.fill.gradient_stops[1].color.rgb = rgb(gradient_to)
        else:
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
            if alpha is not None: _alpha(shp.fill._xPr.find(qn("a:solidFill")), alpha)
        if line is None: shp.line.fill.background()
        else:
            shp.line.color.rgb = rgb(line); shp.line.width = Pt(1)
            if line_alpha is not None: _alpha(shp.line._get_or_add_ln().find(qn("a:solidFill")), line_alpha)
        return shp

    # ── 텍스트 ──
    def text(self, x, y, w, h, content, size, color, bold=False, align="left", anchor="top", spacing=0.0, line_spacing=1.2, max_lines=None, slot=""):
        content = content or ""
        if max_lines:
            est = est_lines(re.sub(r"\*\*|__", "", content), size, w)
            if est > max_lines:
                size -= 1.5; est = est_lines(re.sub(r"\*\*|__", "", content), size, w)
                if est > max_lines:
                    self.ctx.warnings.append({"slide": self.ctx.current_no, "slot": slot, "text": content[:30], "lines": est, "max": max_lines})
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE; tf.vertical_anchor = _ANCHOR[anchor]
        tf.margin_left = tf.margin_right = Inches(0.02); tf.margin_top = tf.margin_bottom = Inches(0.01)
        for i, para in enumerate(content.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = _ALIGN[align]; p.line_spacing = line_spacing; p.space_after = Pt(0)
            for piece in _MARK.split(para):
                if not piece: continue
                r = p.add_run()
                if piece.startswith("**"): r.text = piece[2:-2]; _set_font(r, ds.FONT, size, True, self.accent, spacing)
                elif piece.startswith("__"): r.text = piece[2:-2]; _set_font(r, ds.FONT, size, True, color, spacing)
                else: r.text = piece; _set_font(r, ds.FONT, size, bold, color, spacing)
        return tb

    def kicker(self, text):
        return self.text(ds.MX, ds.KICKER_Y, 12, 0.35, text, ds.SZ["kicker"], self.accent, bold=True, spacing=0.25, max_lines=1, slot="kicker")

    def title(self, text, y=ds.TITLE_Y, size=None, w=ds.CW, max_lines=2):
        return self.text(ds.MX, y, w, 1.70, text, size or ds.SZ["title"], self.ink, bold=True, line_spacing=1.1, max_lines=max_lines, slot="title")

    def lead(self, text, y=ds.LEAD_Y, w=ds.CW):
        return self.text(ds.MX, y, w, 0.80, text, ds.SZ["lead"], self.sub, line_spacing=1.35, max_lines=2, slot="lead")

    def label(self, x, y, w, text, color=None):
        return self.text(x, y, w, 0.3, text, ds.SZ["label"], color or self.accent, bold=True, spacing=0.15, max_lines=1, slot="label")

    def card(self, x, y, w, h, style="plain", top_bar=None):
        st = ds.card_style(self.bg)
        if style == "dark": shp = self.rect(x, y, w, h, fill=ds.BG["dark"], radius=True)
        elif style == "accent": shp = self.rect(x, y, w, h, fill=ds.C["accent"], gradient_to=ds.C["gradient_end"], radius=True)
        elif style == "outline": shp = self.rect(x, y, w, h, fill=None, line=self.accent, radius=True)
        elif style == "tint": shp = self.rect(x, y, w, h, fill=ds.C["accent_tint"], line=ds.C["card_border"], radius=True)
        else: shp = self.rect(x, y, w, h, fill=st["fill"], alpha=st["alpha"], line=st["line"], line_alpha=st["line_alpha"], radius=True)
        if top_bar: self.rect(x, y, w, 0.06, fill=top_bar)
        return shp

    def stat(self, x, y, w, big, unit="", label="", source=None, big_size=None, color=None, label_color=None):
        big_size = big_size or ds.SZ["big"]; color = color or self.ink
        tb = self.text(x, y, w, big_size / 72 * 1.25, big, big_size, color, bold=True, max_lines=1, slot="stat.big")
        if unit:
            p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = " " + unit; _set_font(r, ds.FONT, ds.SZ["unit"], True, color)
        yy = y + big_size / 72 * 1.25 + 0.05
        if label: self.text(x, yy, w, 0.35, label, ds.SZ["body_sm"], self.sub if label_color is None else label_color, max_lines=1, slot="stat.label"); yy += 0.38
        if source: self.text(x, yy, w, 0.3, source, ds.SZ["source"], self.muted, max_lines=1, slot="stat.source")
        return tb

    def body_bullets(self, x, y, w, h, items, size=None, color=None, bullet="·"):
        return self.text(x, y, w, h, "\n".join(f"{bullet} {it}" for it in items), size or ds.SZ["body"], color or self.sub, line_spacing=1.45, slot="bullets")

    def source_line(self, text):
        if text: self.text(ds.MX, ds.SRC_Y, 14.5, 0.3, text, ds.SZ["source"], self.muted, max_lines=1, slot="source_line")

    def page_no(self, n, total):
        self.text(17.5, ds.SRC_Y, 1.375, 0.3, f"{n} / {total}", ds.SZ["page"], self.muted, align="right")

    # ── 이미지 ──
    def _resolve(self, path):
        if not path: return None
        p = Path(path)
        if not p.is_absolute(): p = self.ctx.assets_dir / p
        return p if p.exists() else None

    def picture(self, path, x, y, w, h, cover=True, darken=None, radius=None):
        p = self._resolve(path)
        if not p:
            self.rect(x, y, w, h, fill=ds.C["card_dark"] if self.bg == "dark" else ds.C["card_border"], radius=True)
            self.text(x, y, w, h, "[이미지 확보 필요]", ds.SZ["small"], self.muted, align="center", anchor="middle")
            self.ctx.warnings.append({"slide": self.ctx.current_no, "slot": "image", "text": f"missing {path}", "lines": 0, "max": 0}); return None
        out = self.ctx.tmp / f"s{self.ctx.current_no}_{abs(hash((str(p), x, y, w, h)))}.png"
        px = (int(w * 96), int(h * 96))
        if cover: prep_image.fit_cover(p, out, *px)
        else: prep_image.to_png(p, out)
        if darken: prep_image.darken(out, out, darken)
        if radius: prep_image.round_corners(out, out, radius)
        return self.s.shapes.add_picture(str(out), Inches(x), Inches(y), Inches(w), Inches(h) if cover else None)

    def panel_image(self, path, strength=0.6, x=ds.PANEL_X, w=ds.PANEL_W):
        shp = self.picture(path, x, 0, w, ds.CANVAS_H, cover=True, darken=strength)
        if shp is None: self.rect(x, 0, w, ds.CANVAS_H, fill=ds.C["card_dark"])
        return shp

    def notes(self, text):
        if text: self.s.notes_slide.notes_text_frame.text = text
