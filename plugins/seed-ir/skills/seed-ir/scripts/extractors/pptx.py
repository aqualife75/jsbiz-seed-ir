# -*- coding: utf-8 -*-
from __future__ import annotations
import io, zipfile
from pathlib import Path
from pptx import Presentation
from PIL import Image
from .images import save_image

def extract(path, out_dir) -> dict:
    path = Path(path); out_dir = Path(out_dir)
    prs = Presentation(str(path))
    lines = []
    for i, s in enumerate(prs.slides, 1):
        lines.append(f"\n--- slide {i} ---")
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    if p.text.strip(): lines.append(p.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    lines.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            lines.append("[노트] " + s.notes_slide.notes_text_frame.text.strip())
    images, warnings, k = [], [], 0
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("ppt/media/") and not n.endswith("/"):
                try:
                    im = Image.open(io.BytesIO(z.read(n))); im.load()
                except Exception:  # noqa: BLE001
                    warnings.append(f"skip {n}"); continue
                k += 1
                rec = save_image(im, out_dir, path.stem[:40], k, f"{path.name}:{n}")
                if rec: images.append(rec)
    return {"text": "\n".join(lines), "images": images, "warnings": warnings}
