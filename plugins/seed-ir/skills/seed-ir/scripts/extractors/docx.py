# -*- coding: utf-8 -*-
from __future__ import annotations
import io, zipfile
from pathlib import Path
import docx as _docx
from PIL import Image
from .images import save_image

def extract(path, out_dir) -> dict:
    path = Path(path); out_dir = Path(out_dir)
    d = _docx.Document(str(path))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        lines.append("[표]")
        for row in t.rows:
            lines.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
    images, warnings, k = [], [], 0
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("word/media/") and not n.endswith("/"):
                try:
                    im = Image.open(io.BytesIO(z.read(n))); im.load()
                except Exception:  # noqa: BLE001
                    warnings.append(f"skip {n}"); continue
                k += 1
                rec = save_image(im, out_dir, path.stem[:40], k, f"{path.name}:{n}")
                if rec: images.append(rec)
    return {"text": "\n".join(lines), "images": images, "warnings": warnings}
