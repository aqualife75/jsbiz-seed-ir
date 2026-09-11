# -*- coding: utf-8 -*-
"""deck_spec.json → PPTX. 사용: python build_deck.py spec.json --out deck.pptx"""
from __future__ import annotations
import argparse, json, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import design_system as ds  # noqa: E402
from layouts import LAYOUTS  # noqa: E402
from layouts.base import Canvas, DeckCtx  # noqa: E402

def build(spec_path, out_path) -> dict:
    spec_path = Path(spec_path); spec = json.loads(spec_path.read_text(encoding="utf-8"))
    meta = spec.get("meta", {}); slides = spec["slides"]
    assets = Path(meta.get("assets_dir") or spec_path.parent)
    ctx = DeckCtx(team=meta.get("team", ""), total=len(slides), assets_dir=assets, accent_override=meta.get("accent") if meta.get("accent") not in (None, "", "E0492E") else None)
    prs = Presentation(); prs.slide_width = Inches(ds.CANVAS_W); prs.slide_height = Inches(ds.CANVAS_H)
    blank = prs.slide_layouts[6]
    for sl in slides:
        lay = sl["layout"]
        if lay not in LAYOUTS:
            raise SystemExit(f"unknown layout: {lay} (slide {sl.get('no')})")
        ctx.current_no = sl.get("no", 0)
        cv = Canvas(prs.slides.add_slide(blank), sl.get("background", "cream"), ctx)
        LAYOUTS[lay](cv, sl.get("slots", {}), sl)
    out_path = Path(out_path); out_path.parent.mkdir(parents=True, exist_ok=True); prs.save(str(out_path))
    return {"out": str(out_path), "slides": len(slides), "warnings": ctx.warnings}

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(); ap.add_argument("spec"); ap.add_argument("--out", required=True)
    a = ap.parse_args(argv); rep = build(a.spec, a.out)
    print(f"built {rep['slides']} slides → {rep['out']}  warnings={len(rep['warnings'])}")
    for w in rep["warnings"]: print(f"  ! slide {w['slide']} {w['slot']}: {w['lines']}줄 > {w['max']} '{w['text']}'")
    return 0
if __name__ == "__main__": sys.exit(main())
