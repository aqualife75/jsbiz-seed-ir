# -*- coding: utf-8 -*-
import json
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import build_deck
from layouts import base, LAYOUTS

def _spec(out_dir, img=None):
    return {"meta": {"team": "셀아이", "accent": "E0492E", "assets_dir": str(out_dir)}, "slides": [
        {"no": 1, "layout": "cover", "background": "dark", "slots": {
            "kicker": "TEAM CELLEYE · 사업계획서 2026", "brand": "CellEye", "headline": "라인 위 결함을\n**0.3초**에 잡습니다",
            "subtitle": "이차전지 전극 공정을 위한\nAI 비전 검사 모듈", "tags": ["0.3초 판독", "라인 무정지", "기존 장비 호환"],
            "panel_image": img, "footer_stats": [{"label": "TEAM", "value": "셀아이 · CELLEYE"}, {"label": "TARGET", "value": "전극 라인 400개"},
                                                 {"label": "MODEL", "value": "월 300만 원 구독"}, {"label": "STAGE", "value": "Seed · PoC"}]}, "notes": "표지 대본"},
        {"no": 2, "layout": "statement", "background": "dark", "slots": {"statement": "검사는 자동인데,\n**판독**은 아직 사람입니다", "sub": "문제의 위치"}, "notes": ""},
        {"no": 3, "layout": "trend_cards", "background": "cream", "slots": {
            "title": "라인은 늘고, **불량 비용**은 커지고, 검사 인력은 줍니다", "lead": "세 곡선이 동시에 꺾이는 지점에 공백이 있습니다.",
            "cards": [{"label": "TREND 01 · 라인 증설", "headline": "국내 전극 라인", "big": "400", "unit": "개", "desc": "2년 새 30% 증가했습니다.", "source": "한국배터리산업협회 2025", "tone": "accent"},
                      {"label": "TREND 02 · 불량 비용", "headline": "셀 1개 불량 손실", "big": "12", "unit": "만 원", "desc": "후공정에서 발견 시 3배.", "source": "업계 인터뷰 2026.08", "tone": "amber"},
                      {"label": "TREND 03 · 인력", "headline": "검사원 이직률", "big": "28", "unit": "%", "desc": "야간 판독 인력 부족.", "source": "고용노동부 2025", "tone": "dark"}],
            "band": {"so_what": "판독만 사람이 합니다.", "why_now": ["AI 비전 원가 60% 하락 (2024)", "라인 데이터 표준화", "품질 규제 강화"]},
            "key_gap": "측정은 자동인데 판독은 손"}, "source_line": "출처: 한국배터리산업협회 「전극 라인 현황」 (2025.03) · 수집 2026-09-11", "notes": "트렌드 대본"}]}

def test_build_three_layouts(out_dir):
    img = out_dir / "p.jpg"; Image.new("RGB", (1200, 900), (90, 60, 40)).save(img)
    spec = _spec(out_dir, str(img)); sp = out_dir / "spec.json"; sp.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    rep = build_deck.build(sp, out_dir / "deck.pptx")
    prs = Presentation(str(out_dir / "deck.pptx"))
    assert len(prs.slides) == 3 and prs.slide_width == Inches(20) and prs.slide_height == Inches(11.25)
    texts = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "0.3초" in texts and "3 / 3" in texts and "출처: 한국배터리산업협회" in texts
    assert prs.slides[0].notes_slide.notes_text_frame.text == "표지 대본"
    pics = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]
    assert len(pics) == 1
    fonts = {r.font.name for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.font.name}
    assert fonts == {"Pretendard"}
    assert rep["warnings"] == []

def test_markup_runs(out_dir):
    spec = _spec(out_dir); spec["slides"] = spec["slides"][1:2]
    sp = out_dir / "s.json"; sp.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    build_deck.build(sp, out_dir / "d.pptx")
    prs = Presentation(str(out_dir / "d.pptx"))
    runs = [r for sh in prs.slides[0].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    accent_runs = [r for r in runs if r.text == "판독"]
    assert accent_runs and str(accent_runs[0].font.color.rgb) == "FF6A4D"

def test_overflow_warning(out_dir):
    spec = _spec(out_dir); spec["slides"] = spec["slides"][1:2]
    spec["slides"][0]["slots"]["statement"] = "매우 " * 40
    sp = out_dir / "o.json"; sp.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    rep = build_deck.build(sp, out_dir / "o.pptx")
    assert rep["warnings"] and rep["warnings"][0]["slide"] == 2

def test_registry_has_first_three():
    assert {"cover", "statement", "trend_cards"} <= set(LAYOUTS)
