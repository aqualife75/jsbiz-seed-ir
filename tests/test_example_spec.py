# -*- coding: utf-8 -*-
import json
from pptx import Presentation
import build_deck, validate
from layouts import LAYOUTS

def test_example_spec_covers_all_layouts_and_builds(assets_dir, out_dir):
    spec = json.loads((assets_dir / "spec.example.json").read_text(encoding="utf-8"))
    assert validate.validate_obj(spec, "deck_spec") == []
    used = [s["layout"] for s in spec["slides"]]
    assert set(used) == set(LAYOUTS) and len(LAYOUTS) == 22
    rep = build_deck.build(assets_dir / "spec.example.json", out_dir / "example.pptx")
    assert [w for w in rep["warnings"] if w["slot"] != "image"] == [], rep["warnings"]
    prs = Presentation(str(out_dir / "example.pptx"))
    assert len(prs.slides) == 22 and all(s.notes_slide.notes_text_frame.text for s in prs.slides)
