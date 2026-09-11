# -*- coding: utf-8 -*-
import zipfile
from pathlib import Path
from PIL import Image
import extractors

def _png(path: Path, size=(120, 80), color=(200, 40, 40)):
    Image.new("RGB", size, color).save(path)

def test_docx(out_dir):
    import docx
    d = docx.Document(); d.add_paragraph("문제 정의 MARK1")
    t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "A"; t.rows[0].cells[1].text = "B"
    img = out_dir / "p.png"; _png(img); d.add_picture(str(img))
    src = out_dir / "a.docx"; d.save(src)
    r = extractors.extract_file(src, out_dir / "x")
    assert r["kind"] == "docx" and "MARK1" in r["text"] and "A | B" in r["text"]
    assert len(r["images"]) == 1

def test_pptx(out_dir):
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)); tb.text_frame.text = "시장 규모 MARK2"
    s.notes_slide.notes_text_frame.text = "노트 MARK3"
    img = out_dir / "p.png"; _png(img); s.shapes.add_picture(str(img), Inches(1), Inches(2))
    src = out_dir / "a.pptx"; p.save(src)
    r = extractors.extract_file(src, out_dir / "x")
    assert "MARK2" in r["text"] and "MARK3" in r["text"] and len(r["images"]) == 1

def test_xlsx(out_dir):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "매출"; ws.append(["월", "매출"]); ws.append([1, 1200])
    src = out_dir / "a.xlsx"; wb.save(src)
    r = extractors.extract_file(src, out_dir / "x")
    assert "[시트: 매출]" in r["text"] and "1 | 1200" in r["text"]

def test_pdf(out_dir):
    import fitz
    doc = fitz.open(); page = doc.new_page(); page.insert_text((72, 72), "TAM SAM SOM MARK4")
    img = out_dir / "p.png"; _png(img); page.insert_image(fitz.Rect(72, 100, 272, 250), filename=str(img))
    src = out_dir / "a.pdf"; doc.save(src)
    r = extractors.extract_file(src, out_dir / "x")
    assert "MARK4" in r["text"] and len(r["images"]) >= 1
    assert any(Path(i["file"]).name.startswith("a_page") for i in r["images"])  # 페이지 렌더 포함

def test_hwpx(out_dir):
    src = out_dir / "a.hwpx"
    sec = '<hs:sec xmlns:hp="x"><hp:p><hp:run><hp:t>팀 역량 MARK5</hp:t></hp:run></hp:p><hp:p><hp:run><hp:t>둘째</hp:t></hp:run></hp:p></hs:sec>'
    img = out_dir / "p.png"; _png(img)
    with zipfile.ZipFile(src, "w") as z:
        z.writestr("mimetype", "application/hwp+zip"); z.writestr("Contents/section0.xml", sec)
        z.write(img, "BinData/image1.png")
    r = extractors.extract_file(src, out_dir / "x")
    assert "MARK5\n둘째" in r["text"] and len(r["images"]) == 1

def test_image_and_bmp(out_dir):
    bmp = out_dir / "photo.bmp"; Image.new("RGB", (300, 200), (1, 2, 3)).save(bmp)
    r = extractors.extract_file(bmp, out_dir / "x")
    assert r["kind"] == "image" and r["images"][0]["file"].endswith(".png")

def test_unsupported(out_dir):
    f = out_dir / "a.zzz"; f.write_text("x")
    r = extractors.extract_file(f, out_dir / "x")
    assert r["kind"] == "unsupported" and r["warnings"]
